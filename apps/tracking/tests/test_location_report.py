"""Location Data Report: every GPS record of the selection (paged, sorted and
filtered server-side), per-record data-quality flags, the period analysis
(the Trip Report's own engine), exports carrying ALL records, permissions —
and the removal of the old Operational Reports hub."""

import datetime
import io

import pytest
from django.urls import NoReverseMatch, reverse
from django.utils import timezone

from apps.accounts.models import Permission
from apps.accounts.tests.factories import UserFactory
from apps.clients.tests.factories import ClientFactory
from apps.core.tests.factories import role_with
from apps.core.utils import display_timezone
from apps.tracking import location_report as lr
from apps.tracking import trip_analytics
from apps.tracking.tests.factories import TelemetryEventFactory, TrackingDeviceFactory
from apps.vehicles.tests.factories import VehicleFactory

pytestmark = pytest.mark.django_db

S = datetime.timedelta(seconds=1)
LAT0, LON0, STEP = 12.900000, 77.500000, 0.0009  # ~0.1 km per step


@pytest.fixture(autouse=True)
def _no_map_tiles(settings):
    settings.TRIP_REPORT_MAP_TILE_URL = ""


def _day(offset):
    return timezone.now().astimezone(display_timezone()).date() - datetime.timedelta(days=offset)


def _at(offset_days, hour, minute=0, second=0):
    return datetime.datetime.combine(_day(offset_days), datetime.time(hour, minute, second), tzinfo=display_timezone())


def _vehicle(client_obj=None):
    vehicle = VehicleFactory(client=client_obj)
    TrackingDeviceFactory(vehicle=vehicle)
    return vehicle


def _ev(vehicle, when, *, lat=LAT0, lon=LON0, speed=30, ignition=True, odometer=100.0, metadata=None, client_obj=None):
    return TelemetryEventFactory(
        device=vehicle.tracking_device, vehicle=vehicle, client=client_obj, timestamp=when, ignition=ignition,
        speed=speed, odometer=odometer, latitude=f"{lat:.6f}", longitude=f"{lon:.6f}",
        metadata=metadata if metadata is not None else {"location": "Main Road", "gps_status": 1},
    )


def _drive(vehicle, start, n, every=10 * S, client_obj=None):
    for i in range(n):
        _ev(vehicle, start + i * every, lat=LAT0 + STEP * i, speed=36, odometer=round(100 + 0.1 * i, 1),
            client_obj=client_obj)


def _sel(vehicle, range_key="yesterday", user=None, **kw):
    return lr.resolve_selection(user=user, vehicle=str(vehicle.uuid) if vehicle != "all" else "all",
                                range_key=range_key, **kw)


class TestRecords:
    def test_every_record_is_available_across_pages_oldest_first(self):
        v = _vehicle()
        _drive(v, _at(1, 8), 250)  # a record every 10 s — none reduced to trip start/end
        sel = _sel(v)
        seen = []
        page = lr.records_page(sel, page=1, page_size=100)
        assert (page["total"], page["pages"]) == (250, 3)
        for n in (1, 2, 3):
            seen += [r["ts"] for r in lr.records_page(sel, page=n, page_size=100)["rows"]]
        assert len(seen) == 250 == len(set(seen))
        assert seen == sorted(seen)
        assert sum(1 for _ in lr.iter_records(sel)) == 250

    def test_only_the_selected_vehicle_and_period(self):
        mine, other = _vehicle(), _vehicle()
        _drive(mine, _at(1, 8), 5)
        _drive(other, _at(1, 8), 7)
        _drive(mine, _at(3, 8), 9)  # outside "yesterday"
        rows = lr.records_page(_sel(mine))["rows"]
        assert len(rows) == 5 and {r["registration_number"] for r in rows} == {mine.registration_number}
        assert lr.records_page(_sel("all"))["total"] == 12

    def test_sorting_and_page_size_whitelist(self):
        v = _vehicle()
        for i, speed in enumerate([10, 70, 30, 50]):
            _ev(v, _at(1, 9) + i * 30 * S, lat=LAT0 + STEP * i, speed=speed)
        sel = _sel(v)
        speeds = [r["speed"] for r in lr.records_page(sel, sort="speed", direction="desc")["rows"]]
        assert speeds == [70, 50, 30, 10]
        assert lr.records_page(sel, page_size=12345)["page_size"] == lr.DEFAULT_PAGE_SIZE
        assert lr.records_page(sel, sort="; DROP TABLE x", direction="sideways")["total"] == 4  # whitelisted


class TestQualityFlags:
    def test_problem_records_are_flagged_not_removed(self):
        v = _vehicle()
        t = _at(1, 9)
        _ev(v, t, lat=LAT0)
        _ev(v, t + 10 * S, lat=0, lon=0)  # no fix
        _ev(v, t + 20 * S, lat=LAT0 + STEP)
        # Same instant again: the schema's unique (device, timestamp, lat, lon)
        # rules out exact copies, so real duplicates differ slightly in position.
        _ev(v, t + 20 * S, lat=LAT0 + STEP + 0.00001, odometer=100.1)
        _ev(v, t + 30 * S, lat=LAT0 + 5)  # a 550 km spike in 10 s
        _ev(v, t + 40 * S, lat=LAT0 + 2 * STEP)
        _ev(v, t + 30 * 60 * S, lat=LAT0 + 3 * STEP)  # 30 min gap
        _ev(v, t + 31 * 60 * S, lat=LAT0 + 4 * STEP, speed=-5)  # invalid speed
        _ev(v, t + 32 * 60 * S, lat=LAT0 + 5 * STEP, metadata={"gps_status": "0"})  # device says no fix
        _ev(v, t + 33 * 60 * S, lat=120.0)  # latitude out of range
        sel = _sel(v)
        rows = lr.records_page(sel, page_size=100)["rows"]
        flags = [lr.record_flags(r) for r in rows]
        assert len(rows) == 10  # nothing dropped
        assert flags[1] == ["No GPS fix (0,0)"]
        assert "Duplicate timestamp" in flags[3]
        assert flags[4] == ["GPS jump"] and flags[5] == []  # only the spike, not the fix after it
        assert flags[6] == ["GPS gap before (29 min)"]
        assert flags[7] == ["Invalid speed"]
        assert flags[8] == ["Device reports no GPS fix"]
        assert "Invalid coordinates" in flags[9]
        q = lr.quality_summary(sel)
        assert (q["records"], q["no_fix"], q["duplicate"], q["jump"], q["gap"], q["invalid_speed"],
                q["device_no_fix"], q["invalid_coords"]) == (10, 1, 1, 1, 1, 1, 1, 1)
        assert q["valid_fixes"] == 8
        assert lr.records_page(sel, quality="issues")["total"] == q["flagged"]
        assert lr.records_page(sel, quality="valid")["total"] == 10 - q["flagged"]

    def test_empty_columns_are_reported_so_they_can_be_hidden(self):
        v = _vehicle()
        _drive(v, _at(1, 8), 3)
        q = lr.quality_summary(_sel(v))
        assert q["location"] == 3 and q["gps_status"] == 3
        assert q["altitude"] == 0 and q["satellite_count"] == 0 and q["battery_voltage"] == 0


class TestAnalysis:
    def test_metrics_come_from_the_trip_report_engine(self):
        v = _vehicle()
        t = _at(1, 9)
        for i in range(11):  # 10 x 60 s moving, ~0.1 km apart
            _ev(v, t + i * 60 * S, lat=LAT0 + STEP * i, speed=40)
        for i in range(1, 4):  # 3 min standing, ignition on
            _ev(v, t + (10 + i) * 60 * S, lat=LAT0 + STEP * 10, speed=0)
        _ev(v, t + 14 * 60 * S, lat=LAT0 + STEP * 10, speed=0, ignition=False)
        _ev(v, t + 20 * 60 * S, lat=LAT0 + STEP * 10, speed=0, ignition=False)
        _ev(v, t + 21 * 60 * S, lat=LAT0 + STEP * 10, speed=0, ignition=True)
        sel = _sel(v)
        a = lr.analyse(sel, lr.quality_summary(sel))
        assert a["gps_distance_km"] == pytest.approx(1.0, abs=0.01)
        # The engine classifies each interval by the reading at its START, so the
        # minute 10→11 (40 km/h at minute 10) is still moving.
        assert a["moving_seconds"] == 660
        assert a["idle_seconds"] == 180  # 11→14, ignition on, speed 0
        assert a["engine_off_seconds"] == 420  # 14 → 21, ignition off
        assert a["tracking_seconds"] == 21 * 60
        assert a["ignition_on_seconds"] + a["ignition_off_seconds"] == a["tracking_seconds"]
        assert (a["ignition_off_events"], a["ignition_on_events"]) == (1, 1)
        assert a["max_speed_kmh"] == 40.0
        assert a["avg_moving_speed_kmh"] == pytest.approx(1.0 / (11 / 60), abs=0.1)  # 1 km / 11 min

    def test_too_large_selections_are_refused(self, monkeypatch):
        v = _vehicle()
        _drive(v, _at(1, 8), 5)
        monkeypatch.setattr(trip_analytics, "EXPORT_MAX_GPS_RECORDS", 3)
        sel = _sel(v)
        with pytest.raises(trip_analytics.ReportError) as excinfo:
            lr.analyse(sel, lr.quality_summary(sel))
        assert excinfo.value.status == 413


class TestSelection:
    def test_vehicle_is_required_and_ranges_are_validated(self):
        with pytest.raises(trip_analytics.ReportError, match="Select a vehicle"):
            lr.resolve_selection(user=None, vehicle="", range_key="today")
        v = _vehicle()
        with pytest.raises(trip_analytics.ReportError, match="on or before"):
            lr.resolve_selection(user=None, vehicle=str(v.uuid), range_key="custom", from_str="2026-02-10",
                                 to_str="2026-02-01")

    @pytest.mark.parametrize("range_key, days", [("today", 1), ("yesterday", 1), ("last3", 3), ("last5", 5),
                                                 ("last7", 7)])
    def test_date_presets(self, range_key, days):
        sel = _sel(_vehicle(), range_key)
        assert (sel.end_date - sel.start_date).days + 1 == days


def _viewer(client, client_obj=None):
    client.force_login(UserFactory(role=role_with((Permission.Module.TRACKING_DEVICE, Permission.Action.VIEW)),
                                   client=client_obj))


class TestApi:
    def test_data_endpoint_shape(self, client):
        v = _vehicle()
        _drive(v, _at(1, 8), 120)
        _viewer(client)
        body = client.get(reverse("location-report-data"),
                          {"vehicle": str(v.uuid), "range": "yesterday", "analysis": "1"}).json()
        assert body["records"]["total"] == 120 and len(body["records"]["rows"]) == 100
        assert body["quality"]["records"] == 120 and body["analysis"]["gps_distance_km"] > 0
        assert "altitude" not in body["columns"] and "location" in body["columns"]
        row = body["records"]["rows"][0]
        for key in ("timestamp", "latitude", "longitude", "speed", "ignition", "odometer", "gps_status", "flags"):
            assert key in row
        page2 = client.get(reverse("location-report-data"),
                           {"vehicle": str(v.uuid), "range": "yesterday", "page": 2}).json()
        assert len(page2["records"]["rows"]) == 20 and "analysis" not in page2  # page turns skip the analysis

    def test_permissions_and_client_scoping(self, client):
        mine, theirs = ClientFactory(), ClientFactory()
        other = _vehicle(theirs)
        _drive(other, _at(1, 8), 3, client_obj=theirs)
        client.force_login(UserFactory(role=None))
        assert client.get(reverse("location-report-data"), {"vehicle": "all"}).status_code == 403
        assert client.get(reverse("tracking:location_report")).status_code == 403
        _viewer(client, mine)
        response = client.get(reverse("location-report-data"), {"vehicle": str(other.uuid), "range": "yesterday"})
        assert response.status_code == 404
        assert client.get(reverse("location-report-export"),
                          {"vehicle": str(other.uuid), "range": "yesterday", "type": "xlsx"}).status_code == 404

    def test_excel_export_holds_every_record_not_just_one_page(self, client):
        from openpyxl import load_workbook

        v = _vehicle()
        _drive(v, _at(1, 8), 250)
        _ev(v, _at(1, 9), lat=0, lon=0)  # one flagged record
        _viewer(client)
        response = client.get(reverse("location-report-export"),
                              {"vehicle": str(v.uuid), "range": "yesterday", "type": "xlsx"})
        assert response.status_code == 200
        wb = load_workbook(io.BytesIO(response.content))
        assert wb.sheetnames == ["Summary", "Location Data", "Data Quality"]
        data = wb["Location Data"]
        assert data.max_row == 1 + 251
        header = [c.value for c in data[1]]
        assert {"GPS timestamp", "Latitude", "Longitude", "Speed (km/h)", "Ignition", "Odometer (km)",
                "Data status"} <= set(header)
        assert "Altitude (m)" not in header  # empty for this selection
        assert wb["Data Quality"].max_row == 1 + 1

    def test_pdf_export_and_its_limit(self, client, monkeypatch):
        v = _vehicle()
        _drive(v, _at(1, 8), 30)
        _viewer(client)
        params = {"vehicle": str(v.uuid), "range": "yesterday", "type": "pdf"}
        ok = client.get(reverse("location-report-export"), params)
        assert ok.status_code == 200 and ok.content.startswith(b"%PDF")
        assert "location-data_" in ok["Content-Disposition"]
        monkeypatch.setattr(lr, "PDF_MAX_ROWS", 10)
        refused = client.get(reverse("location-report-export"), params)
        assert refused.status_code == 413 and "Download Excel" in refused.json()["detail"]

    def test_no_records_and_missing_vehicle_are_readable_errors(self, client):
        v = _vehicle()
        _viewer(client)
        empty = client.get(reverse("location-report-export"), {"vehicle": str(v.uuid), "range": "today", "type": "pdf"})
        assert empty.status_code == 404 and "No GPS records" in empty.json()["detail"]
        assert client.get(reverse("location-report-data"), {"range": "today"}).status_code == 400


def test_page_and_menu(client):
    _viewer(client)
    response = client.get(reverse("tracking:location_report"))
    html = response.content.decode()
    assert response.context["active_nav_section"] == "Reports"
    assert response.context["active_nav_label"] == "Location Data Report"
    for pill in ("today", "yesterday", "last3", "last5", "last7", "custom"):
        assert f'data-range="{pill}"' in html
    assert 'id="locVehicle"' in html and 'id="locApplyBtn"' in html and 'data-report-download="xlsx"' in html
    assert 'id="locCustomApply"' not in html  # this page applies everything with one button


def test_operational_reports_hub_is_gone(client):
    _viewer(client)
    with pytest.raises(NoReverseMatch):
        reverse("reports:report_hub")
    assert client.get("/reports/").status_code == 404
    assert "Operational Reports" not in client.get(reverse("core:dashboard")).content.decode()
    assert not Permission.objects.filter(module="report").exists()
