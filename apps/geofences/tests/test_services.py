"""Geofence monitoring (apps.geofences.services) through the real ingestion
path: ENTRY / EXIT / SPEED_LIMIT alerts for ASSIGNED vehicles, one per state
change per (vehicle, geofence), with boundary hysteresis and crossing
confirmation; plus the Alert Report, exports and bell for them."""

import datetime
import io
from decimal import Decimal

import pytest
from django.utils import timezone

from apps.accounts.models import Permission
from apps.accounts.tests.factories import UserFactory
from apps.alerts.models import Alert
from apps.clients.tests.factories import ClientFactory
from apps.core.tests.factories import role_with
from apps.geofences import services
from apps.geofences.models import Geofence, GeofenceEvent
from apps.geofences.tests.factories import GeofenceFactory
from apps.notifications.models import Notification
from apps.tracking.providers.base import NormalizedEvent
from apps.tracking.services import TelemetryIngestionService
from apps.tracking.tests.factories import TrackingDeviceFactory
from apps.vehicles.tests.factories import VehicleFactory

pytestmark = pytest.mark.django_db

S = datetime.timedelta(seconds=1)
CENTER = (12.971600, 77.594600)
DEG_PER_M = 1 / 111_320  # latitude degrees per metre


def _vehicle(**kwargs):
    vehicle = VehicleFactory(client=kwargs.pop("client", None) or ClientFactory(), **kwargs)
    TrackingDeviceFactory(vehicle=vehicle)
    return vehicle


def _fence(*vehicles, **kwargs):
    defaults = {"center_latitude": f"{CENTER[0]:.6f}", "center_longitude": f"{CENTER[1]:.6f}",
                "radius_meters": 500, "geofence_type": Geofence.GeofenceType.ENTRY}
    defaults.update(kwargs)
    fence = GeofenceFactory(**defaults)
    fence.vehicles.set(vehicles)
    return fence


def _at(metres_north, when, speed=30, **metadata):
    """A reading ``metres_north`` of the geofence center."""
    lat = CENTER[0] + metres_north * DEG_PER_M
    return NormalizedEvent(timestamp=when, latitude=Decimal(f"{lat:.6f}"), longitude=Decimal(f"{CENTER[1]:.6f}"),
                           speed=Decimal(str(speed)), ignition=True,
                           metadata={"location": "MG Road, Bengaluru", "gps_status": 1, **metadata})


def _drive(vehicle, metres, *, speeds=None, start=None, step=10 * S, batch=3):
    start = start or timezone.now() - datetime.timedelta(minutes=30)
    readings = [_at(m, start + i * step, speed=(speeds[i] if speeds else 30)) for i, m in enumerate(metres)]
    for i in range(0, len(readings), batch):
        TelemetryIngestionService.ingest_events(device=vehicle.tracking_device, events=readings[i:i + batch])
    return readings


def _alerts(vehicle, category):
    return list(Alert.objects.filter(vehicle=vehicle, category=category).order_by("occurred_at"))


@pytest.fixture
def watcher():
    return UserFactory(role=role_with((Permission.Module.ALERT, Permission.Action.VIEW)))


class TestGeometry:
    def test_same_point_is_zero_distance(self):
        assert services.distance_meters("12.9716", "77.5946", "12.9716", "77.5946") == 0

    def test_known_distance_is_approximately_correct(self):
        assert 110_000 < services.distance_meters(12.0, 77.0, 13.0, 77.0) < 112_000

    def test_circle_inside_outside(self):
        fence = GeofenceFactory(center_latitude="12.971600", center_longitude="77.594600", radius_meters=500)
        assert services.outside_by(fence, CENTER[0] + 400 * DEG_PER_M, CENTER[1]) < 0
        assert services.outside_by(fence, CENTER[0] + 600 * DEG_PER_M, CENTER[1]) > 0

    def test_polygon_inside_outside_and_edge_distance(self):
        square = [[12.97, 77.59], [12.97, 77.60], [12.98, 77.60], [12.98, 77.59]]
        fence = GeofenceFactory(shape=Geofence.Shape.POLYGON, polygon=square, center_latitude="12.975000",
                                center_longitude="77.595000", radius_meters=800)
        assert services.outside_by(fence, 12.975, 77.595) < 0  # center
        outside = services.outside_by(fence, 12.975, 77.6005)  # ~54 m east of the edge
        assert 40 < outside < 70
        assert services.outside_by(fence, 13.5, 77.595) > 50_000

    def test_hysteresis_leaving_needs_the_tolerance(self, settings):
        settings.GEOFENCE_EXIT_TOLERANCE_METERS = 20
        fence = GeofenceFactory(center_latitude="12.971600", center_longitude="77.594600", radius_meters=500)
        just_out = CENTER[0] + 510 * DEG_PER_M
        assert services.is_inside(fence, just_out, CENTER[1], currently_inside=True)  # 10 m out: still inside
        assert not services.is_inside(fence, just_out, CENTER[1], currently_inside=False)  # but not an entry


class TestEntry:
    def test_entering_raises_one_medium_entry_alert(self, watcher):
        v = _vehicle()
        fence = _fence(v, name="Bangalore City")
        readings = _drive(v, [900, 800, 700, 400, 300, 200, 100, 50, 0])
        [alert] = _alerts(v, Alert.Category.GEOFENCE_ENTRY)
        assert (alert.get_category_display(), alert.get_severity_display()) == ("Geofence Entry", "Medium")
        assert alert.content_object == fence and alert.message == f"Vehicle {v.registration_number} entered Bangalore City."
        assert alert.occurred_at == readings[3].timestamp  # the first reading inside
        assert alert.latitude == readings[3].latitude and alert.location == "MG Road, Bengaluru"
        assert Notification.objects.filter(alert=alert, recipient=watcher).count() == 1
        assert GeofenceEvent.objects.filter(geofence=fence, vehicle=v, event_type="ENTER").count() == 1

    def test_staying_inside_is_not_another_entry(self, watcher):
        v = _vehicle()
        _fence(v)
        _drive(v, [900, 400, 300] + [100] * 30)
        assert len(_alerts(v, Alert.Category.GEOFENCE_ENTRY)) == 1 and Notification.objects.count() == 1

    def test_exit_then_reenter_is_a_second_entry_and_closes_the_first(self, watcher):
        v = _vehicle()
        _fence(v)
        _drive(v, [900, 400, 300, 200, 700, 800, 900, 300, 200])
        first, second = _alerts(v, Alert.Category.GEOFENCE_ENTRY)
        assert first.signal_cleared_at is not None and second.signal_cleared_at is None
        assert first.signal_cleared_at <= second.occurred_at

    def test_entry_geofence_never_raises_exit_alerts(self, watcher):
        v = _vehicle()
        _fence(v)
        _drive(v, [900, 400, 300, 700, 800])
        assert _alerts(v, Alert.Category.GEOFENCE_EXIT) == []


class TestExit:
    def test_leaving_raises_one_exit_alert(self, watcher):
        v = _vehicle()
        fence = _fence(v, name="Depot", geofence_type=Geofence.GeofenceType.EXIT)
        readings = _drive(v, [0, 100, 200, 700, 800, 900, 1000])
        [alert] = _alerts(v, Alert.Category.GEOFENCE_EXIT)
        assert (alert.get_category_display(), alert.get_severity_display()) == ("Geofence Exit", "Medium")
        assert alert.occurred_at == readings[3].timestamp and alert.message == f"Vehicle {v.registration_number} exited Depot."
        assert alert.content_object == fence
        assert _alerts(v, Alert.Category.GEOFENCE_ENTRY) == []  # an EXIT geofence alerts on exit only
        assert Notification.objects.count() == 1

    def test_staying_outside_is_not_another_exit(self, watcher):
        v = _vehicle()
        _fence(v, geofence_type=Geofence.GeofenceType.EXIT)
        _drive(v, [0, 100, 700, 800] + [2000] * 20)
        assert len(_alerts(v, Alert.Category.GEOFENCE_EXIT)) == 1

    def test_never_inside_never_exits(self, watcher):
        v = _vehicle()
        _fence(v, geofence_type=Geofence.GeofenceType.EXIT)
        _drive(v, [2000, 1900, 1800, 1700])
        assert _alerts(v, Alert.Category.GEOFENCE_EXIT) == []


class TestSpeedLimit:
    def _speed_fence(self, v, limit=40):
        return _fence(v, name="School Zone", geofence_type=Geofence.GeofenceType.SPEED_LIMIT, speed_limit_kmh=limit,
                      radius_meters=2000)

    def test_spec_sequence(self, watcher):
        """35, 40 no alert; 41 NEW; 45, 48, 52, 55 same; 38 cleared; 45 NEW."""
        v = _vehicle()
        fence = self._speed_fence(v)
        speeds = [35, 35, 40, 41, 45, 48, 52, 55, 38, 45]
        _drive(v, [100] * len(speeds), speeds=speeds)
        first, second = _alerts(v, Alert.Category.GEOFENCE_SPEEDING)
        assert (first.get_category_display(), first.get_severity_display()) == ("Geofence Speeding", "High")
        assert first.speed_limit == 40 and first.speed == Decimal("55.00")  # top speed of the episode
        assert first.signal_cleared_at is not None and second.signal_cleared_at is None
        assert first.content_object == fence
        assert first.message == "Speeding in School Zone: 41 km/h (limit 40 km/h)."
        assert Notification.objects.filter(alert__category="GEOFENCE_SPEEDING").count() == 2

    def test_speed_outside_the_geofence_is_ignored(self, watcher):
        v = _vehicle()
        self._speed_fence(v, limit=30)
        _drive(v, [5000, 4800, 4600, 4400], speeds=[60, 60, 60, 60])
        assert _alerts(v, Alert.Category.GEOFENCE_SPEEDING) == []

    def test_leaving_the_zone_stops_monitoring_and_clears(self, watcher):
        v = _vehicle()
        self._speed_fence(v, limit=30)
        _drive(v, [100, 100, 100, 2500, 2600, 2700], speeds=[35, 40, 45, 60, 60, 60])
        [alert] = _alerts(v, Alert.Category.GEOFENCE_SPEEDING)
        assert alert.signal_cleared_at is not None and alert.speed == Decimal("45.00")  # not the 60s outside

    def test_entering_already_fast_alerts_once(self, watcher):
        v = _vehicle()
        self._speed_fence(v, limit=30)
        _drive(v, [3000, 2900, 1500, 1400, 1300], speeds=[50, 50, 50, 50, 50])
        assert len(_alerts(v, Alert.Category.GEOFENCE_SPEEDING)) == 1


class TestRobustness:
    def test_gps_jitter_on_the_boundary_does_not_flap(self, watcher):
        v = _vehicle()
        _fence(v)
        _drive(v, [900, 400, 300] + [495, 505, 498, 510, 502, 499, 515] * 3)
        assert len(_alerts(v, Alert.Category.GEOFENCE_ENTRY)) == 1
        assert GeofenceEvent.objects.filter(vehicle=v).count() == 1  # never "exited"

    def test_a_single_gps_jump_is_not_an_entry(self, watcher):
        v = _vehicle()
        _fence(v)
        _drive(v, [3000, 3000, 0, 3000, 3000])
        assert _alerts(v, Alert.Category.GEOFENCE_ENTRY) == []

    def test_no_fix_readings_are_not_evaluated(self, watcher):
        v = _vehicle()
        _fence(v)
        start = timezone.now() - datetime.timedelta(minutes=20)
        readings = [NormalizedEvent(timestamp=start + i * 10 * S, latitude=Decimal("0"), longitude=Decimal("0"),
                                    speed=Decimal("0"), ignition=True, metadata={}) for i in range(5)]
        readings += [_at(0, start + 60 * S, gps_status="0"), _at(0, start + 70 * S, gps_status="0")]
        TelemetryIngestionService.ingest_events(device=v.tracking_device, events=readings)
        assert _alerts(v, Alert.Category.GEOFENCE_ENTRY) == []

    def test_unassigned_vehicles_and_inactive_geofences_are_not_evaluated(self, watcher):
        assigned, other = _vehicle(), _vehicle()
        _fence(assigned)
        _fence(assigned, status=Geofence.Status.INACTIVE)
        _drive(other, [900, 400, 300, 200])
        _drive(assigned, [900, 400, 300, 200])
        assert _alerts(other, Alert.Category.GEOFENCE_ENTRY) == []
        assert len(_alerts(assigned, Alert.Category.GEOFENCE_ENTRY)) == 1  # the active one only

    def test_state_is_per_vehicle_and_geofence(self, watcher):
        a, b = _vehicle(), _vehicle()
        f1 = _fence(a, b)
        f2 = _fence(a, center_latitude=f"{CENTER[0] + 3000 * DEG_PER_M:.6f}")
        _drive(a, [900, 400, 300])
        _drive(b, [5000, 4800])
        assert [x.content_object for x in _alerts(a, Alert.Category.GEOFENCE_ENTRY)] == [f1]
        assert _alerts(b, Alert.Category.GEOFENCE_ENTRY) == []
        _drive(a, [2900, 3000, 3000], start=timezone.now() - datetime.timedelta(minutes=10))
        assert {x.content_object for x in _alerts(a, Alert.Category.GEOFENCE_ENTRY)} == {f1, f2}

    def test_redelivered_and_out_of_order_batches_change_nothing(self, watcher):
        v = _vehicle()
        _fence(v)
        start = timezone.now() - datetime.timedelta(minutes=20)
        readings = [_at(m, start + i * 10 * S) for i, m in enumerate([900, 800, 400, 300, 200])]
        for batch in ([readings[3]], readings[:2], readings[2:], readings, readings[4:]):
            TelemetryIngestionService.ingest_events(device=v.tracking_device, events=batch)
        [alert] = _alerts(v, Alert.Category.GEOFENCE_ENTRY)
        assert alert.occurred_at == readings[2].timestamp
        assert Notification.objects.count() == 1 and GeofenceEvent.objects.filter(vehicle=v).count() == 1

    def test_polygon_geofence_end_to_end(self, watcher):
        v = _vehicle()
        lat0, lon0 = CENTER
        d = 500 * DEG_PER_M
        _fence(v, shape=Geofence.Shape.POLYGON, radius_meters=710,
               polygon=[[lat0 - d, lon0 - d], [lat0 - d, lon0 + d], [lat0 + d, lon0 + d], [lat0 + d, lon0 - d]])
        _drive(v, [900, 800, 400, 300, 0])
        assert len(_alerts(v, Alert.Category.GEOFENCE_ENTRY)) == 1

    def test_geofences_never_touch_trips_or_other_alerts(self, watcher):
        v = _vehicle()
        _fence(v, geofence_type=Geofence.GeofenceType.SPEED_LIMIT, speed_limit_kmh=40, radius_meters=2000)
        _drive(v, [100, 100, 100], speeds=[50, 50, 50])
        kinds = set(Alert.objects.filter(vehicle=v).values_list("category", flat=True))
        assert kinds == {"GEOFENCE_SPEEDING"}


class TestPermissionsAndNotification:
    def test_only_users_who_can_see_the_vehicle_are_notified(self):
        a, b = ClientFactory(), ClientFactory()
        view = role_with((Permission.Module.ALERT, Permission.Action.VIEW))
        user_a, user_b, staff = UserFactory(role=view, client=a), UserFactory(role=view, client=b), UserFactory(role=view)
        v = _vehicle(client=a)
        _fence(v)
        _drive(v, [900, 400, 300])
        assert set(Notification.objects.values_list("recipient_id", flat=True)) == {user_a.pk, staff.pk}
        assert not Notification.objects.filter(recipient=user_b).exists()

    def test_bell_announces_it(self, client, watcher):
        v = _vehicle()
        _fence(v, name="Bangalore City")
        _drive(v, [900, 400, 300])
        client.force_login(watcher)
        [announced] = client.get("/notifications/feed/").json()["alerts"]
        assert (announced["type_label"], announced["level_label"]) == ("Geofence Entry", "Medium")
        assert announced["message"] == f"Vehicle {v.registration_number} entered Bangalore City."
        assert client.get("/notifications/pulse/").json()["latest_alert"] == announced["notification_uuid"]


class TestReportAndExports:
    def test_report_filters_and_columns(self, client):
        from openpyxl import load_workbook

        a, b = ClientFactory(), ClientFactory()
        va, vb = _vehicle(client=a), _vehicle(client=b)
        zone = _fence(va, name="Bangalore City", geofence_type=Geofence.GeofenceType.SPEED_LIMIT, speed_limit_kmh=40,
                      radius_meters=2000)
        depot = _fence(vb, name="Bravo Depot")
        _drive(va, [100, 100, 100], speeds=[52, 50, 30])
        _drive(vb, [900, 400, 300])
        client.force_login(UserFactory(role=role_with((Permission.Module.ALERT, Permission.Action.VIEW))))
        rows = client.get("/api/v1/alerts/report/", {"geofence": str(zone.uuid)}).json()["alerts"]["rows"]
        assert [(r["type_label"], r["level_label"], r["geofence"]["name"], r["speed_limit"], r["speed"])
                for r in rows] == [("Geofence Speeding", "High", "Bangalore City", 40, 52.0)]
        assert {r["type"] for r in client.get("/api/v1/alerts/report/", {"alert_type": "GEOFENCE_ENTRY"}).json()[
            "alerts"]["rows"]} == {"GEOFENCE_ENTRY"}
        xlsx = client.get("/api/v1/alerts/report/export/", {"type": "xlsx"})
        header, *data = list(load_workbook(io.BytesIO(xlsx.content))["Alerts"].iter_rows(values_only=True))
        got = {(dict(zip(header, r))["Alert"], dict(zip(header, r))["Geofence"], dict(zip(header, r))["Speed limit (km/h)"])
               for r in data}
        assert got == {("Geofence Speeding", "Bangalore City", 40), ("Geofence Entry", "Bravo Depot", None)}
        pdf = client.get("/api/v1/alerts/report/export/", {"type": "pdf", "geofence": str(depot.uuid)})
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")

    def test_client_users_never_see_other_clients_geofence_alerts_or_names(self, client):
        a, b = ClientFactory(), ClientFactory()
        va, vb = _vehicle(client=a), _vehicle(client=b)
        _fence(va, name="Alpha Zone")
        bravo = _fence(vb, name="Bravo Secret Zone")
        _drive(va, [900, 400, 300])
        _drive(vb, [900, 400, 300])
        client.force_login(UserFactory(role=role_with((Permission.Module.ALERT, Permission.Action.VIEW)), client=a))
        rows = client.get("/api/v1/alerts/report/").json()["alerts"]["rows"]
        assert [r["geofence"]["name"] for r in rows] == ["Alpha Zone"]
        assert "Bravo Secret Zone" not in client.get("/alerts/report/").content.decode()
        assert client.get("/api/v1/alerts/report/", {"geofence": str(bravo.uuid)}).status_code == 404


class TestEntryAndExit:
    """One ENTRY_AND_EXIT geofence: an Entry alert on every entry AND an Exit alert on every exit."""

    def _both(self, *vehicles, **kwargs):
        return _fence(*vehicles, name="Bangalore City", geofence_type=Geofence.GeofenceType.ENTRY_AND_EXIT, **kwargs)

    def test_spec_table(self, watcher):
        """1-2 OUTSIDE nothing; 3 INSIDE ENTRY; 4-6 nothing; 7 OUTSIDE EXIT; 8-9 nothing; 10 INSIDE ENTRY.
        Each crossing is confirmed by the next reading (GEOFENCE_CONFIRM_READINGS = 2)."""
        v = _vehicle(registration_number="KA01AB1234")
        fence = self._both(v)
        metres = [900, 900, 300, 300, 300, 300, 900, 900, 900, 300, 300]
        readings = _drive(v, metres)
        entries = _alerts(v, Alert.Category.GEOFENCE_ENTRY)
        exits = _alerts(v, Alert.Category.GEOFENCE_EXIT)
        assert [a.occurred_at for a in entries] == [readings[2].timestamp, readings[9].timestamp]
        assert [a.occurred_at for a in exits] == [readings[6].timestamp]
        assert {a.content_object for a in entries + exits} == {fence}  # one geofence, both event types
        assert Geofence.objects.count() == 1
        assert [a.get_severity_display() for a in entries + exits] == ["Medium"] * 3
        assert entries[0].message == "Vehicle KA01AB1234 entered Bangalore City."
        assert exits[0].message == "Vehicle KA01AB1234 exited Bangalore City."
        # One notification per event, with the event in its title (bell, popup and sound once each).
        titles = list(Notification.objects.filter(recipient=watcher).order_by("created_at").values_list("title", flat=True))
        assert titles == ["Geofence Entry alert — KA01AB1234", "Geofence Exit alert — KA01AB1234",
                          "Geofence Entry alert — KA01AB1234"]
        # Entry alert closes when the vehicle leaves; Exit alert closes when it comes back.
        assert entries[0].signal_cleared_at == readings[6].timestamp
        assert exits[0].signal_cleared_at == readings[9].timestamp

    def test_staying_inside_or_outside_never_repeats(self, watcher):
        v = _vehicle()
        self._both(v)
        _drive(v, [900] * 10 + [300] * 20 + [900] * 20)
        assert len(_alerts(v, Alert.Category.GEOFENCE_ENTRY)) == 1 and len(_alerts(v, Alert.Category.GEOFENCE_EXIT)) == 1
        assert Notification.objects.count() == 2

    def test_boundary_jitter_creates_no_false_alerts(self, watcher):
        v = _vehicle()
        self._both(v)
        _drive(v, [900, 300, 300] + [495, 505, 498, 510, 502, 499, 515, 490] * 3)
        assert len(_alerts(v, Alert.Category.GEOFENCE_ENTRY)) == 1 and _alerts(v, Alert.Category.GEOFENCE_EXIT) == []

    def test_permissions_and_report(self, client):
        from openpyxl import load_workbook

        a, b = ClientFactory(), ClientFactory()
        va, vb = _vehicle(client=a), _vehicle(client=b)
        self._both(va, vb)
        user_a = UserFactory(role=role_with((Permission.Module.ALERT, Permission.Action.VIEW)), client=a)
        _drive(va, [900, 300, 300, 900, 900])
        _drive(vb, [900, 300, 300])
        assert set(Notification.objects.filter(recipient=user_a).values_list("alert__vehicle", flat=True)) == {va.pk}
        client.force_login(user_a)
        rows = client.get("/api/v1/alerts/report/").json()["alerts"]["rows"]
        assert sorted((r["type_label"], r["level_label"], r["geofence"]["name"]) for r in rows) == [
            ("Geofence Entry", "Medium", "Bangalore City"), ("Geofence Exit", "Medium", "Bangalore City")]
        xlsx = client.get("/api/v1/alerts/report/export/", {"type": "xlsx"})
        header, *data = list(load_workbook(io.BytesIO(xlsx.content))["Alerts"].iter_rows(values_only=True))
        assert sorted(dict(zip(header, r))["Alert"] for r in data) == ["Geofence Entry", "Geofence Exit"]
        assert client.get("/api/v1/alerts/report/export/", {"type": "pdf"}).content.startswith(b"%PDF")

    def test_form_and_live_map_data(self, client):
        from django.urls import reverse

        client.force_login(UserFactory(role=role_with(
            (Permission.Module.GEOFENCE, Permission.Action.VIEW), (Permission.Module.GEOFENCE, Permission.Action.CREATE),
            (Permission.Module.TRACKING_DEVICE, Permission.Action.VIEW))))
        html = client.get(reverse("geofences:geofence_create")).content.decode()
        for label in ("Entry", "Exit", "Speed Limit", "Entry and Exit"):
            assert f"<span>{label}</span>" in html
        assert 'type="radio" name="geofence_type"' in html
        response = client.post(reverse("geofences:geofence_create"), {
            "code": "BOTH01", "name": "Bangalore City", "status": "ACTIVE", "geofence_type": "ENTRY_AND_EXIT",
            "speed_limit_kmh": "40", "shape": "CIRCLE", "center_latitude": "12.9716", "center_longitude": "77.5946",
            "radius_meters": 3000})
        assert response.status_code == 302
        fence = Geofence.objects.get(code="BOTH01")
        assert fence.speed_limit_kmh is None and fence.alerts_on_entry and fence.alerts_on_exit
        [item] = client.get(reverse("geofences:geofence_map_data")).json()["results"]
        assert (item["type"], item["type_name"], item["speed_limit_kmh"]) == ("ENTRY_AND_EXIT", "Entry and Exit", None)
        live = client.get(reverse("tracking:live_map")).content.decode()
        assert 'data-geofence-count="ENTRY_AND_EXIT"' in live
