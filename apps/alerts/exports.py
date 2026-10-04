"""PDF / Excel for the Alert Report — every alert of the applied selection
(apps.alerts.report.iter_alerts: the same filtered, scoped Alert rows as the
page), on the Trip Report's shared export blocks (apps.tracking.trip_exports)."""

from django.utils import timezone

from apps.alerts import report as ar
from apps.tracking import trip_analytics as ta
from apps.tracking.trip_exports import (
    DASH,
    FMT_COORD,
    FMT_DATE,
    FMT_DT,
    FMT_DUR,
    FMT_INT,
    FMT_TIME,
    _Pdf,
    _summary_sheet,
    _xl_dt,
    _Xlsx,
    fmt_dt,
)


def _span(selection):
    s, e = selection.start_date, selection.end_date
    return f"{s:%d %b %Y}" if s == e else f"{s:%d %b %Y} – {e:%d %b %Y}"


def _meta(selection):
    return ta._meta(
        title="Alert Report", subtitle=f"{selection.vehicle_label} · {_span(selection)}",
        period_start=selection.period_start, period_end=min(selection.period_end, timezone.now()),
        period_label=f"{ta.RANGE_LABELS.get(selection.range_key, '')} ({_span(selection)})",
        filters_label=(f"Vehicle: {selection.vehicle_label} · Alert: {selection.type_label}"
                       f" · Level: {selection.level_label} · Status: {selection.status_label}"
                       + (f" · Geofence: {selection.geofence.name}" if selection.geofence else "")),
        now=timezone.now(), tz=selection.tz, vehicles=[],
    )


def _summary_rows(totals):
    rows = [("Alerts", f"{totals['matching']:,}")]
    rows += [(f"{label} alerts", f"{totals['by_type'].get(key, 0):,}") for key, label in ar.TYPE_CHOICES]
    rows += [(f"Level: {label}", f"{n:,}") for key, label in ar.LEVEL_CHOICES
             if (n := totals["by_level"].get(key, 0))]
    rows += [
        ("New (not yet acknowledged)", f"{totals['open']:,}"),
        ("Acknowledged", f"{totals['acknowledged']:,}"),
        ("Resolved", f"{totals['resolved']:,}"),
        ("Still active (panic held / still idling)", f"{totals['signal_active']:,}"),
        ("Vehicles with alerts", f"{totals['vehicles']:,}"),
    ]
    return rows


def _status_text(alert):
    return ar.STATUS_LABELS.get(alert.status, alert.get_status_display())


def _num(value, places, unit=""):
    return f"{float(value):,.{places}f}{unit}" if value is not None else DASH


def alert_report_pdf(selection):
    from reportlab.platypus import Spacer

    totals = ar.summary(selection)
    if totals["matching"] > ar.PDF_MAX_ROWS:
        raise ta.ReportError(
            f"This selection has {totals['matching']:,} alerts — too many for a readable PDF (maximum "
            f"{ar.PDF_MAX_ROWS:,}). Download Excel for the complete list, or choose a shorter period.", status=413)
    tz = selection.tz
    pdf = _Pdf(_meta(selection))
    story = pdf.title_block()
    story += [pdf.p(
        "Each row is one alert event, from the reading that started it (date, time, location, speed). Panic "
        "(critical): a continuous panic signal from the device (panic input ≥ 10 V) counts once. Idle (medium): "
        "ignition ON and the vehicle stationary for at least the vehicle's idle threshold (default 5 min); "
        "Duration is the real idle time, until it moved or the ignition went off. Status is the alert's handling "
        "(New → Acknowledged → Resolved).", "note"), Spacer(1, 6)]
    cards = [("Alerts", f"{totals['matching']:,}", selection.type_label)]
    cards += [(label, f"{totals['by_type'].get(key, 0):,}", "events") for key, label in ar.TYPE_CHOICES]
    cards += [
        ("New", f"{totals['open']:,}", "not yet acknowledged"),
        ("Acknowledged", f"{totals['acknowledged']:,}", ""),
        ("Resolved", f"{totals['resolved']:,}", ""),
        ("Vehicles", f"{totals['vehicles']:,}", "with alerts"),
    ]
    story += [pdf.section("1. Summary"), pdf.kpi_grid(cards, cols=4)]

    headers = ["#", "Date", "Time", "Vehicle", "Driver", "Alert", "Level", "Duration", "Geofence", "Location",
               "Speed", "Voltage", "Status"]
    widths = [0.35, 0.8, 0.6, 0.95, 0.95, 0.75, 0.6, 0.65, 0.9, 2.3, 0.85, 0.55, 0.8]
    rows = []
    for i, alert in enumerate(ar.iter_alerts(selection), start=1):
        location = alert.location or (
            f"{alert.latitude}, {alert.longitude}" if alert.latitude is not None else DASH)
        rows.append([
            i, fmt_dt(alert.occurred_at, tz, "%d %b %Y"), fmt_dt(alert.occurred_at, tz, "%H:%M:%S"),
            alert.vehicle.registration_number if alert.vehicle_id else DASH,
            alert.driver.get_full_name() if alert.driver_id else DASH,
            alert.get_category_display(), alert.get_severity_display(),
            ar.duration_text(ar.duration_seconds(alert)) or DASH,
            (g.name if (g := ar.alert_geofence(alert)) else DASH), location,
            _num(alert.speed, 1, " km/h") + (f" (limit {alert.speed_limit})" if alert.speed_limit else ""),
            _num(alert.voltage, 2, " V"), _status_text(alert),
        ])
    story.append(pdf.section("2. Alerts", f"{len(rows):,} alert(s), oldest first."))
    if rows:
        story.append(pdf.table(headers, rows, widths, align_right=[0, 7, 10, 11], wrap=[5, 8, 9], dense=True))
    else:
        story.append(pdf.p("No alerts match this selection."))
    return pdf.build(story)


# Excel ----------------------------------------------------------------------

_XL_COLUMNS = [
    # header, number format, width, value(alert, tz)
    ("Alert ID", None, 38, lambda a, tz: str(a.uuid)),
    ("Date", FMT_DATE, 12, lambda a, tz: _xl_dt(a.occurred_at, tz).date() if a.occurred_at else None),
    ("Time", FMT_TIME, 10, lambda a, tz: _xl_dt(a.occurred_at, tz).time() if a.occurred_at else None),
    ("Vehicle", None, 14, lambda a, tz: a.vehicle.registration_number if a.vehicle_id else None),
    ("Vehicle ID", None, 14, lambda a, tz: (a.vehicle.vehicle_code or None) if a.vehicle_id else None),
    ("Client", None, 22, lambda a, tz: a.client.client_name if a.client_id else None),
    ("Driver", None, 20, lambda a, tz: a.driver.get_full_name() if a.driver_id else None),
    ("Alert", None, 12, lambda a, tz: a.get_category_display()),
    ("Level", None, 10, lambda a, tz: a.get_severity_display()),
    ("Start time", FMT_DT, 20, lambda a, tz: _xl_dt(a.occurred_at, tz) if a.occurred_at else None),
    ("Alert time", FMT_DT, 20, lambda a, tz: _xl_dt(a.triggered_at or a.occurred_at, tz) if a.occurred_at else None),
    ("End time", FMT_DT, 20, lambda a, tz: _xl_dt(a.signal_cleared_at, tz) if a.signal_cleared_at else None),
    ("Duration", FMT_DUR, 11, lambda a, tz: (ar.duration_seconds(a) / 86400) if ar.duration_seconds(a) is not None else None),
    ("Duration (s)", FMT_INT, 11, lambda a, tz: ar.duration_seconds(a)),
    ("Latitude", FMT_COORD, 12, lambda a, tz: float(a.latitude) if a.latitude is not None else None),
    ("Longitude", FMT_COORD, 12, lambda a, tz: float(a.longitude) if a.longitude is not None else None),
    ("Geofence", None, 22, lambda a, tz: g.name if (g := ar.alert_geofence(a)) else None),
    ("Location", None, 60, lambda a, tz: a.location or None),
    ("Speed (km/h)", "0.0", 11, lambda a, tz: float(a.speed) if a.speed is not None else None),
    ("Speed limit (km/h)", FMT_INT, 12, lambda a, tz: a.speed_limit),
    ("Voltage (V)", "0.00", 11, lambda a, tz: float(a.voltage) if a.voltage is not None else None),
    ("Ignition", None, 9, lambda a, tz: {True: "On", False: "Off"}.get(a.ignition)),
    ("Odometer (km)", "#,##0.0", 13, lambda a, tz: float(a.odometer) if a.odometer is not None else None),
    ("Status", None, 13, lambda a, tz: _status_text(a)),
    ("Acknowledged by", None, 18, lambda a, tz: a.acknowledged_by.get_full_name() if a.acknowledged_by_id else None),
    ("Acknowledged at", FMT_DT, 20, lambda a, tz: _xl_dt(a.acknowledged_at, tz) if a.acknowledged_at else None),
    ("Resolved by", None, 18, lambda a, tz: a.resolved_by.get_full_name() if a.resolved_by_id else None),
    ("Resolved at", FMT_DT, 20, lambda a, tz: _xl_dt(a.resolved_at, tz) if a.resolved_at else None),
    ("Created at", FMT_DT, 20, lambda a, tz: _xl_dt(a.created_at, tz)),
]


def alert_report_xlsx(selection):
    from openpyxl.utils import get_column_letter

    totals = ar.summary(selection)
    x = _Xlsx(_meta(selection))
    tz = selection.tz
    notes = [
        "Alerts sheet: one row per alert event of the selection, oldest first. A continuous panic signal is one "
        "event; an idle period (ignition ON, stationary for the vehicle's idle threshold) is one event. Position, "
        "speed and voltage are those of the reading that started it.",
        "Start time = when the condition began; Alert time = when it became an alert (idle: threshold reached); "
        "End time = panic released / vehicle moved or ignition off (empty = still active); Duration = start → end. "
        "Status = the alert's handling (New / Acknowledged / Resolved). Times are in the report timezone.",
    ]
    _summary_sheet(x, "Summary", [(label, value, None) for label, value in _summary_rows(totals)], [], notes)

    ws = x.wb.create_sheet("Alerts")
    for j, (_header, _fmt, width, _value) in enumerate(_XL_COLUMNS):
        ws.column_dimensions[get_column_letter(j + 1)].width = width
    ws.freeze_panes = "B2"
    ws.append([x.cell(ws, header, font=x.header_font, fill=x.header_fill,
                      align=x.Alignment(vertical="center", wrap_text=True)) for header, _f, _w, _v in _XL_COLUMNS])
    count = 0
    for count, alert in enumerate(ar.iter_alerts(selection), start=1):
        if count > ar.EXCEL_MAX_ROWS:
            raise ta.ReportError("Too many alerts for one Excel sheet — please choose a shorter period.", status=413)
        ws.append([x.cell(ws, v, fmt=fmt) if (v := value(alert, tz)) is not None else None
                   for _h, fmt, _w, value in _XL_COLUMNS])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_XL_COLUMNS))}{max(count, 1) + 1}"
    return x.save()

