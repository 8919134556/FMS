"""DEVICE BATTERY alerts: comms ``device_battery_voltage`` (= batteryvoltage /
1000, V) -> exact bands (0-<2 V disconnected HIGH, 2-<3 V low MEDIUM, >=3 V
normal) -> one alert per state change, on the same level engine as Panic and
Main Power but with its own flags. The voltage is an alert input only: it must
never reach Live Tracking (page, feed, current position) or a visible history
field."""

import datetime
import io
from decimal import Decimal

import pytest
from django.utils import timezone

from apps.accounts.models import Permission
from apps.accounts.tests.factories import UserFactory
from apps.alerts.events import device_battery_flags, device_battery_state, main_power_flags
from apps.alerts.models import Alert
from apps.clients.tests.factories import ClientFactory
from apps.core.tests.factories import role_with
from apps.notifications.models import Notification
from apps.tracking import comms_sync
from apps.tracking.models import TelemetryEvent, VehicleCurrentTelemetry
from apps.tracking.providers.base import NormalizedEvent
from apps.tracking.services import TelemetryIngestionService
from apps.tracking.tests.factories import TrackingDeviceFactory
from apps.vehicles.tests.factories import VehicleFactory

pytestmark = pytest.mark.django_db

S = datetime.timedelta(seconds=1)
CUT, LOW = Alert.Category.DEVICE_BATTERY_DISCONNECTED, Alert.Category.DEVICE_BATTERY_LOW_VOLTAGE
IMEI = "350424067315497"


def _base():
    return timezone.now() - datetime.timedelta(minutes=30)


def _vehicle(**kwargs):
    vehicle = VehicleFactory(client=kwargs.pop("client", None) or ClientFactory(), **kwargs)
    TrackingDeviceFactory(vehicle=vehicle)
    return vehicle


def _r(when, battery, i=0, mainpower=12.6):
    meta = {"location": "Bengaluru", "mainpower": mainpower, **main_power_flags(mainpower)}
    if battery is not None:
        meta["device_battery_voltage"] = battery
        meta.update(device_battery_flags(battery))
    return NormalizedEvent(timestamp=when, latitude=Decimal(f"{12.98 + i * 0.00001:.6f}"), longitude=Decimal("77.53"),
                           speed=Decimal("0"), ignition=True, metadata=meta)


def _series(vehicle, start, volts, *, mainpower=12.6, step=30 * S):
    readings = [_r(start + k * step, v, i=k, mainpower=mainpower) for k, v in enumerate(volts)]
    for k in range(0, len(readings), 3):
        TelemetryIngestionService.ingest_events(device=vehicle.tracking_device, events=readings[k:k + 3])
    return readings


def _alerts(vehicle, category):
    return list(Alert.objects.filter(category=category, vehicle=vehicle).order_by("occurred_at"))


@pytest.fixture
def watcher():
    return UserFactory(role=role_with((Permission.Module.ALERT, Permission.Action.VIEW)))


class TestRanges:
    @pytest.mark.parametrize("volts, expected", [
        (0, "DISCONNECTED"), (0.00, "DISCONNECTED"), (1.99, "DISCONNECTED"), ("1.999", "DISCONNECTED"),
        (2.00, "LOW"), (2.50, "LOW"), (2.99, "LOW"), ("2.999", "LOW"),
        (3.00, "NORMAL"), (3.01, "NORMAL"), (3.95, "NORMAL"), (4.096, "NORMAL"), (5.00, "NORMAL"),
        (None, None), ("", None), ("x", None), (-0.5, None), (float("nan"), None),
    ])
    def test_exact_bands(self, volts, expected):
        assert device_battery_state(volts) == expected

    @pytest.mark.parametrize("millivolts, expected", [
        (0, "DISCONNECTED"), (1999, "DISCONNECTED"), (2000, "LOW"), (2999, "LOW"), (3000, "NORMAL"),
        (3999, "NORMAL"), (4000, "NORMAL"), (4096, "NORMAL"),
    ])
    def test_raw_batteryvoltage_through_the_comms_conversion(self, millivolts, expected):
        """comms stores device_battery_voltage = batteryvoltage / 1000; the bridge classifies that value."""
        tz = datetime.timezone(datetime.timedelta(hours=5, minutes=30))
        row = {"id": 1, "unitno": IMEI, "tracktime": datetime.datetime(2026, 10, 4, 16, 29, 24), "lat": 12.98,
               "lon": 77.53, "speed": 0.0, "direction": "0", "ignition": 1, "odometer": 1000, "gpsodometer": 0,
               "gpsstatus": 1, "location": "Bengaluru", "device_battery_voltage": millivolts / 1000.0}
        meta = comms_sync.row_to_event(row, tz).metadata
        assert device_battery_state(meta["device_battery_voltage"]) == expected
        assert meta["battery_cut"] == int(expected == "DISCONNECTED") and meta["battery_low"] == int(expected == "LOW")

    def test_null_battery_has_no_flags(self):
        assert device_battery_flags(None) == {}


class TestStateMachine:
    def test_low_voltage_sequence_from_the_spec(self, watcher):
        """4.2, 4.1 normal; 2.8 NEW Low; 2.7, 2.5, 2.2 same; 4.2 recovered."""
        v = _vehicle()
        t = _base()
        _series(v, t, [4.2, 4.1, 2.8, 2.7, 2.5, 2.2, 4.2])
        [low] = _alerts(v, LOW)
        assert (low.get_category_display(), low.get_severity_display()) == ("Device Battery Low Voltage", "Medium")
        assert low.occurred_at == t + 2 * 30 * S and low.signal_cleared_at == t + 6 * 30 * S
        assert low.voltage == Decimal("2.80") and _alerts(v, CUT) == []
        assert Notification.objects.filter(alert=low).count() == 1

    def test_disconnected_sequence_then_a_new_low(self, watcher):
        """4.2; 1.9 NEW Disconnected; 1.5, 1.0, 0 same; 4.2 recovered; 4.2; 2.7 NEW Low."""
        v = _vehicle()
        t = _base()
        _series(v, t, [4.2, 1.9, 1.5, 1.0, 0, 4.2, 4.2, 2.7])
        [cut] = _alerts(v, CUT)
        [low] = _alerts(v, LOW)
        assert (cut.get_category_display(), cut.get_severity_display()) == ("Device Battery Disconnected", "High")
        assert cut.occurred_at == t + 30 * S and cut.signal_cleared_at == t + 5 * 30 * S
        assert low.occurred_at == t + 7 * 30 * S and low.signal_cleared_at is None
        assert Notification.objects.count() == 2

    def test_low_to_disconnected_to_low_to_normal(self, watcher):
        v = _vehicle()
        t = _base()
        _series(v, t, [4.1, 2.5, 1.5, 2.5, 4.1])
        assert [(a.occurred_at, a.signal_cleared_at) for a in _alerts(v, LOW)] == [
            (t + 30 * S, t + 60 * S), (t + 90 * S, t + 120 * S)]
        assert [(a.occurred_at, a.signal_cleared_at) for a in _alerts(v, CUT)] == [(t + 60 * S, t + 90 * S)]
        assert Notification.objects.count() == 3  # only state changes INTO an alert state

    @pytest.mark.parametrize("volts, expected", [
        (0.00, CUT), (1.99, CUT), (2.00, LOW), (2.99, LOW), (3.00, None), (3.01, None), (3.95, None),
        (5.00, None),
    ])
    def test_boundaries_end_to_end(self, watcher, volts, expected):
        v = _vehicle()
        _series(v, _base(), [4.1, volts])
        raised = list(Alert.objects.filter(vehicle=v).values_list("category", flat=True))
        assert raised == ([expected] if expected else [])

    def test_null_and_negative_never_alert_nor_clear(self, watcher):
        v = _vehicle()
        t = _base()
        _series(v, t, [4.1, None, -1, None, 4.1])
        assert Alert.objects.filter(vehicle=v).count() == 0
        _series(v, t + 6 * 30 * S, [1.0, None, -1, 4.1])
        [cut] = _alerts(v, CUT)
        assert cut.signal_cleared_at == t + 9 * 30 * S

    def test_redelivery_never_repeats(self, watcher):
        v = _vehicle()
        readings = _series(v, _base(), [4.1, 1.0, 1.0])
        for _ in range(3):
            TelemetryIngestionService.ingest_events(device=v.tracking_device, events=[readings[-1]])
        assert len(_alerts(v, CUT)) == 1 and Notification.objects.count() == 1

    def test_vehicles_are_independent(self, watcher):
        a, b, c = _vehicle(), _vehicle(), _vehicle()
        t = _base()
        _series(a, t, [4.1, 1.5])
        _series(b, t, [4.1, 2.5])
        _series(c, t, [4.1, 4.2])
        assert (len(_alerts(a, CUT)), len(_alerts(a, LOW))) == (1, 0)
        assert (len(_alerts(b, CUT)), len(_alerts(b, LOW))) == (0, 1)
        assert Alert.objects.filter(vehicle=c).count() == 0


class TestSeparateFromMainPower:
    def test_battery_alerts_do_not_touch_main_power_and_vice_versa(self, watcher):
        v = _vehicle()
        t = _base()
        _series(v, t, [4.1, 1.0, 1.0], mainpower=12.6)  # battery cut, main power normal
        assert Alert.objects.filter(vehicle=v, category__in=["MAIN_POWER_DISCONNECTED", "LOW_VOLTAGE"]).count() == 0
        _series(v, t + 3 * 30 * S, [1.0, 1.0], mainpower=0)  # main power cut while battery stays cut
        assert len(_alerts(v, "MAIN_POWER_DISCONNECTED")) == 1
        assert len(_alerts(v, CUT)) == 1  # still the same battery episode
        _series(v, t + 5 * 30 * S, [4.1], mainpower=0)  # battery recovers, main power still cut
        assert _alerts(v, CUT)[0].signal_cleared_at is not None
        assert _alerts(v, "MAIN_POWER_DISCONNECTED")[0].signal_cleared_at is None

    def test_the_bridge_reads_each_supply_from_its_own_column(self):
        tz = datetime.timezone(datetime.timedelta(hours=5, minutes=30))
        row = {"id": 1, "unitno": IMEI, "tracktime": datetime.datetime(2026, 10, 4, 16, 29, 24), "lat": 12.98,
               "lon": 77.53, "speed": 0.0, "direction": "0", "ignition": 1, "odometer": 1000, "gpsodometer": 0,
               "gpsstatus": 1, "location": "Bengaluru", "mainpower": 12.79, "device_battery_voltage": 0.0}
        event = comms_sync.row_to_event(row, tz)
        assert (event.metadata["power_cut"], event.metadata["battery_cut"]) == (0, 1)
        assert event.external_power is True  # main power, not the battery
        assert event.battery_voltage is None  # the battery voltage is not a visible history field


class TestNotification:
    def test_notification_and_bell_show_the_alert_not_the_voltage(self, client, watcher):
        v = _vehicle(registration_number="KA01AB1234")
        _series(v, _base(), [4.1, 1.5])
        n = Notification.objects.get(recipient=watcher)
        assert n.title == "Device Battery Disconnected alert — KA01AB1234"
        assert "Level: High" in n.body and "Device battery disconnected." in n.body and "1.5" not in n.body
        client.force_login(watcher)
        [announced] = client.get("/notifications/feed/").json()["alerts"]
        assert (announced["type_label"], announced["level_label"]) == ("Device Battery Disconnected", "High")
        assert "1.5" not in str(announced)
        pulse = client.get("/notifications/pulse/").json()
        assert pulse["latest_alert"] == str(n.uuid)


class TestReportAndExports:
    def test_report_filter_pdf_and_excel(self, client):
        from openpyxl import load_workbook

        v = _vehicle()
        _series(v, timezone.now() - datetime.timedelta(minutes=20), [4.1, 2.5, 1.5, 4.1])
        client.force_login(UserFactory(role=role_with((Permission.Module.ALERT, Permission.Action.VIEW))))
        rows = client.get("/api/v1/alerts/report/", {"alert_type": CUT}).json()["alerts"]["rows"]
        assert [(r["type_label"], r["level_label"]) for r in rows] == [("Device Battery Disconnected", "High")]
        rows = client.get("/api/v1/alerts/report/", {"alert_type": LOW}).json()["alerts"]["rows"]
        assert [(r["type_label"], r["level_label"]) for r in rows] == [("Device Battery Low Voltage", "Medium")]
        xlsx = client.get("/api/v1/alerts/report/export/", {"type": "xlsx"})
        header, *data = list(load_workbook(io.BytesIO(xlsx.content))["Alerts"].iter_rows(values_only=True))
        pairs = {(dict(zip(header, r))["Alert"], dict(zip(header, r))["Level"]) for r in data}
        assert pairs == {("Device Battery Low Voltage", "Medium"), ("Device Battery Disconnected", "High")}
        pdf = client.get("/api/v1/alerts/report/export/", {"type": "pdf", "alert_type": CUT})
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


class TestNotOnLiveTracking:
    def test_never_on_the_live_feed_page_or_current_position(self, client):
        v = _vehicle()
        _series(v, _base(), [2.5])
        client.force_login(UserFactory(role=role_with((Permission.Module.TRACKING_DEVICE, Permission.Action.VIEW))))
        [item] = client.get("/api/v1/tracking/fleet/current/").json()["results"]
        assert not any("batter" in key for key in item)
        assert "2.5" not in str({k: val for k, val in item.items() if k not in ("latitude", "longitude")})
        html = client.get("/tracking/live/").content.decode().lower()
        assert "device battery" not in html and "battery voltage" not in html
        assert not any("batter" in f.name for f in VehicleCurrentTelemetry._meta.get_fields())
        assert TelemetryEvent.objects.get(vehicle=v).battery_voltage is None

    def test_not_on_the_location_data_report_either(self, client):
        v = _vehicle()
        _series(v, _base(), [2.5])
        client.force_login(UserFactory(role=role_with((Permission.Module.TRACKING_DEVICE, Permission.Action.VIEW))))
        body = client.get("/api/v1/tracking/location-report/", {"vehicle": str(v.uuid), "range": "today",
                                                               "analysis": "1"}).json()
        assert "battery_voltage" not in body.get("columns", [])


class TestCommsBridgeEndToEnd:
    def test_comms_rows_raise_one_alert_per_state_change(self, watcher):
        v = _vehicle()
        v.tracking_device.imei = IMEI
        v.tracking_device.save(update_fields=["imei"])
        now_local = timezone.now().astimezone(comms_sync._local_timezone()).replace(tzinfo=None, microsecond=0)
        t = now_local - datetime.timedelta(minutes=5)
        volts = [4.05, 3.952, 2.5, 0.0, 0.0, 4.05]  # 3.952 V is normal now; 2.5 low; 0 disconnected
        rows = [{"id": k + 1, "unitno": IMEI, "tracktime": t + k * 30 * S, "lat": 12.98, "lon": 77.53 + k * 1e-5,
                 "speed": 0.0, "direction": "0", "ignition": 1, "odometer": 1000, "gpsodometer": 0, "gpsstatus": 1,
                 "location": "Bengaluru", "mainpower": 12.7, "device_battery_voltage": bv}
                for k, bv in enumerate(volts)]
        comms_sync.apply_rows(rows, comms_sync.SyncReport())
        assert len(_alerts(v, LOW)) == 1 and len(_alerts(v, CUT)) == 1
        assert Notification.objects.filter(recipient=watcher).count() == 2
