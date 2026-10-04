"""Alert Report (page, API, PDF/Excel), its client scoping, the in-place
acknowledge/resolve, and the bell feed's alert announcements — all reading the
same Alert events raised by apps.alerts.events."""

import datetime
import io
from decimal import Decimal

import pytest
from django.core.management import call_command
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import Permission, Role
from apps.accounts.tests.factories import UserFactory
from apps.alerts.models import Alert
from apps.alerts.tests.factories import AlertFactory
from apps.audit.models import AuditLog
from apps.clients.tests.factories import ClientFactory
from apps.core.permissions import user_has_permission
from apps.core.tests.factories import role_with
from apps.core.utils import display_timezone
from apps.notifications.services import notify
from apps.tracking.tests.factories import TrackingDeviceFactory
from apps.vehicles.tests.factories import VehicleFactory

pytestmark = pytest.mark.django_db

M, A = Permission.Module, Permission.Action
REPORT_API = "/api/v1/alerts/report/"
EXPORT_API = "/api/v1/alerts/report/export/"


def _today_at(hour, minute=0):
    local = timezone.now().astimezone(display_timezone())
    when = local.replace(hour=hour, minute=minute, second=0, microsecond=0)
    return min(when, timezone.now() - datetime.timedelta(minutes=1))


def _panic(vehicle, when, **kwargs):
    defaults = {
        "dedupe_key": f"PANIC:{vehicle.pk}:{when.isoformat()}", "category": Alert.Category.PANIC,
        "severity": Alert.Severity.CRITICAL, "title": f"Panic alert — {vehicle.registration_number}",
        "vehicle": vehicle, "client": vehicle.client, "occurred_at": when, "last_signal_at": when,
        "latitude": Decimal("12.985287"), "longitude": Decimal("77.539518"), "location": "Bengaluru",
        "speed": Decimal("0"), "ignition": True, "voltage": Decimal("12.37"), "odometer": Decimal("108605.1"),
    }
    defaults.update(kwargs)
    return Alert.objects.create(**defaults)


@pytest.fixture
def world():
    a, b = ClientFactory(client_name="Client Alpha"), ClientFactory(client_name="Client Bravo")
    va = VehicleFactory(client=a, registration_number="ALPHA-001")
    vb = VehicleFactory(client=b, registration_number="BRAVO-001")
    for v in (va, vb):
        TrackingDeviceFactory(vehicle=v)
    view = role_with((M.ALERT, A.VIEW))
    return {
        "a": a, "b": b, "va": va, "vb": vb,
        "staff": UserFactory(role=role_with((M.ALERT, A.VIEW), (M.ALERT, A.UPDATE))),
        "user_a": UserFactory(role=view, client=a), "user_b": UserFactory(role=view, client=b),
        "alert_a": _panic(va, _today_at(9)), "alert_b": _panic(vb, _today_at(10)),
    }


def _get(client, user, url, **params):
    client.force_login(user)
    return client.get(url, params)


class TestPage:
    def test_report_page_needs_alert_view(self, client, world):
        assert _get(client, world["staff"], reverse("alerts:alert_report")).status_code == 200
        nobody = UserFactory(role=role_with((M.VEHICLE, A.VIEW)))
        assert _get(client, nobody, reverse("alerts:alert_report")).status_code == 403
        assert _get(client, nobody, REPORT_API).status_code == 403

    def test_vehicle_filter_lists_only_what_the_user_can_see(self, client, world):
        html = _get(client, world["user_a"], reverse("alerts:alert_report")).content.decode()
        assert "ALPHA-001" in html and "BRAVO-001" not in html

    def test_menu_entry_opens_the_alert_report(self, client, world):
        response = _get(client, world["user_a"], reverse("alerts:alert_report"))
        urls = {item["url"] for section in response.context["nav_sections"] for item in section["items"]}
        assert reverse("alerts:alert_report") in urls


class TestScoping:
    def test_staff_sees_every_client(self, client, world):
        body = _get(client, world["staff"], REPORT_API).json()
        assert {r["registration_number"] for r in body["alerts"]["rows"]} == {"ALPHA-001", "BRAVO-001"}

    def test_a_client_user_sees_only_their_own_alerts(self, client, world):
        body = _get(client, world["user_a"], REPORT_API).json()
        assert [r["registration_number"] for r in body["alerts"]["rows"]] == ["ALPHA-001"]
        assert body["summary"]["total"] == 1
        assert body["can_update"] is False

    def test_another_clients_vehicle_or_alert_is_not_found(self, client, world):
        assert _get(client, world["user_a"], REPORT_API, vehicle=str(world["vb"].uuid)).status_code == 404
        assert _get(client, world["user_a"], f"/api/v1/alerts/{world['alert_b'].uuid}/").status_code == 404
        assert _get(client, world["user_a"], f"/api/v1/alerts/{world['alert_a'].uuid}/").status_code == 200

    def test_exports_are_scoped_too(self, client, world):
        from openpyxl import load_workbook

        response = _get(client, world["user_b"], EXPORT_API, type="xlsx")
        rows = list(load_workbook(io.BytesIO(response.content)).worksheets[1].iter_rows(values_only=True))
        assert [r[3] for r in rows[1:]] == ["BRAVO-001"]

    def test_history_stays_with_the_client_who_owned_the_vehicle(self, client, world):
        """The vehicle moves to Bravo later: Alpha keeps its past alert, Bravo doesn't get it."""
        world["va"].client = world["b"]
        world["va"].save(update_fields=["client"])
        assert _get(client, world["user_a"], REPORT_API).json()["summary"]["total"] == 1
        assert {r["registration_number"] for r in _get(client, world["user_b"], REPORT_API).json()["alerts"]["rows"]} == {
            "BRAVO-001"}

    def test_alerts_work_list_is_scoped_and_hides_operational_alerts_from_clients(self, client, world):
        AlertFactory(title="Trip TRP-1 is delayed")
        html = _get(client, world["user_a"], reverse("alerts:alert_list")).content.decode()
        assert "ALPHA-001" in html and "BRAVO-001" not in html and "Trip TRP-1 is delayed" not in html
        assert "Trip TRP-1 is delayed" in _get(client, world["staff"], reverse("alerts:alert_list")).content.decode()

    def test_client_users_can_view_but_never_act(self, client, world):
        assert user_has_permission(world["user_a"], "alert", "view")
        assert not user_has_permission(world["user_a"], "alert", "update")
        client.force_login(world["user_a"])
        response = client.post(reverse("alerts:alert_acknowledge", kwargs={"uuid": world["alert_a"].uuid}))
        assert response.status_code == 403


class TestFilters:
    def test_vehicle_type_status_and_period(self, client, world):
        _panic(world["va"], _today_at(9, 30), status=Alert.Status.RESOLVED)
        _panic(world["va"], timezone.now() - datetime.timedelta(days=2))
        staff = world["staff"]
        assert _get(client, staff, REPORT_API).json()["summary"]["total"] == 3  # today
        assert _get(client, staff, REPORT_API, range="last3").json()["summary"]["total"] == 4
        only_a = _get(client, staff, REPORT_API, vehicle=str(world["va"].uuid)).json()
        assert only_a["summary"]["total"] == 2 and {r["registration_number"] for r in only_a["alerts"]["rows"]} == {
            "ALPHA-001"}
        resolved = _get(client, staff, REPORT_API, status="RESOLVED").json()
        assert resolved["alerts"]["total"] == 1 and resolved["summary"]["total"] == 3  # pills keep their counts
        assert resolved["summary"]["resolved"] == 1 and resolved["summary"]["open"] == 2
        assert _get(client, staff, REPORT_API, alert_type="PANIC").json()["summary"]["by_type"]["PANIC"] == 3
        assert _get(client, staff, REPORT_API, alert_type="SPEEDING").status_code == 400

    def test_operational_alerts_are_not_in_the_vehicle_report(self, client, world):
        AlertFactory(category=Alert.Category.DOCUMENT_EXPIRED)
        assert _get(client, world["staff"], REPORT_API).json()["summary"]["total"] == 2

    def test_paging_and_sorting_happen_on_the_server(self, client, world):
        for i in range(30):
            _panic(world["va"], _today_at(8) - datetime.timedelta(minutes=i + 1), voltage=Decimal(f"{10 + i / 10:.2f}"))
        staff = world["staff"]
        page2 = _get(client, staff, REPORT_API, page=2, page_size=25).json()["alerts"]
        assert (page2["total"], page2["pages"], page2["page"], len(page2["rows"])) == (32, 2, 2, 7)
        first = _get(client, staff, REPORT_API, sort="voltage", dir="asc", page_size=25).json()["alerts"]["rows"]
        assert first[0]["voltage"] == 10.0
        newest = _get(client, staff, REPORT_API).json()["alerts"]["rows"][0]
        assert newest["registration_number"] == "BRAVO-001"  # default: newest first

    def test_bad_custom_range_is_a_readable_error(self, client, world):
        response = _get(client, world["staff"], REPORT_API, range="custom", **{"from": "2026-10-05", "to": "2026-10-01"})
        assert response.status_code == 400 and "start date" in response.json()["detail"]


class TestDetailAndWorkflow:
    def test_detail_has_everything_the_popup_shows(self, client, world):
        alert = world["alert_a"]
        alert.signal_cleared_at = alert.occurred_at + datetime.timedelta(seconds=5)
        alert.save()
        body = _get(client, world["staff"], f"/api/v1/alerts/{alert.uuid}/").json()
        row = body["alert"]
        assert row["type_label"] == "Panic" and row["status_label"] == "New"
        assert row["registration_number"] == "ALPHA-001" and row["client"] == "Client Alpha"
        assert (row["latitude"], row["voltage"], row["ignition"], row["odometer"]) == (12.985287, 12.37, True, 108605.1)
        assert row["signal_active"] is False and row["duration_seconds"] == 5
        assert body["can_update"] is True

    def test_acknowledge_then_resolve_in_place(self, client, world):
        client.force_login(world["staff"])
        alert = world["alert_a"]
        json_headers = {"HTTP_ACCEPT": "application/json"}
        response = client.post(reverse("alerts:alert_acknowledge", kwargs={"uuid": alert.uuid}), **json_headers)
        assert response.json() == {"ok": True, "status": "ACKNOWLEDGED"}
        response = client.post(reverse("alerts:alert_resolve", kwargs={"uuid": alert.uuid}), **json_headers)
        assert response.json()["status"] == "RESOLVED"
        alert.refresh_from_db()
        assert alert.acknowledged_by == world["staff"] and alert.resolved_by == world["staff"]
        assert AuditLog.objects.filter(module="alert", entity_id=str(alert.pk)).count() == 2
        row = _get(client, world["staff"], f"/api/v1/alerts/{alert.uuid}/").json()["alert"]
        assert row["status_label"] == "Resolved" and row["resolved_by"]


class TestExports:
    def test_pdf_and_excel_hold_the_filtered_alerts(self, client, world):
        from openpyxl import load_workbook

        staff = world["staff"]
        pdf = _get(client, staff, EXPORT_API, type="pdf", vehicle=str(world["va"].uuid))
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
        assert 'filename="alert-report_ALPHA-001_' in pdf["Content-Disposition"]
        xlsx = _get(client, staff, EXPORT_API, type="xlsx")
        book = load_workbook(io.BytesIO(xlsx.content))
        assert book.sheetnames == ["Summary", "Alerts"]
        header, *rows = list(book["Alerts"].iter_rows(values_only=True))
        assert header[:9] == ("Alert ID", "Date", "Time", "Vehicle", "Vehicle ID", "Client", "Driver", "Alert type", "Severity")
        assert [r[3] for r in rows] == ["ALPHA-001", "BRAVO-001"]  # oldest first
        assert {r[0] for r in rows} == {str(world["alert_a"].uuid), str(world["alert_b"].uuid)}
        voltage = header.index("Voltage (V)")
        assert rows[0][voltage] == 12.37
        assert AuditLog.objects.filter(module="alert", entity="AlertReport", action=AuditLog.Action.EXPORT).count() == 2

    def test_nothing_to_export_is_a_message_not_an_empty_file(self, client, world):
        response = _get(client, world["staff"], EXPORT_API, type="pdf", range="yesterday")
        assert response.status_code == 404 and "No alerts" in response.json()["detail"]
        assert _get(client, world["staff"], EXPORT_API, type="csv").status_code == 400

    def test_pdf_continues_across_pages(self, client, world):
        for i in range(80):
            _panic(world["va"], _today_at(8) - datetime.timedelta(minutes=i + 1))
        pdf = _get(client, world["staff"], EXPORT_API, type="pdf").content
        pages = pdf.count(b"/Type /Page") - pdf.count(b"/Type /Pages")
        assert pages >= 3  # 82 rows: the table (with its header row) continues over several pages


class TestBellFeed:
    def test_feed_announces_unread_alert_notifications_once_read_they_stop(self, client, world):
        staff = world["staff"]
        alert = world["alert_a"]
        n = notify(staff, title="Panic alert — ALPHA-001", alert=alert, link_url=alert.link_url or "/alerts/report/")
        notify(staff, title="Ordinary message")
        client.force_login(staff)
        body = client.get(reverse("notifications:notification_feed")).json()
        assert [x["alert_uuid"] for x in body["alerts"]] == [str(alert.uuid)]
        announced = body["alerts"][0]
        assert announced["vehicle"] == "ALPHA-001" and announced["type_label"] == "Panic"
        assert announced["location"] == "Bengaluru" and announced["notification_uuid"] == str(n.uuid)
        assert {r["title"]: r["is_alert"] for r in body["results"]} == {
            "Panic alert — ALPHA-001": True, "Ordinary message": False}
        client.post(reverse("notifications:notification_mark_read", kwargs={"uuid": n.uuid}), HTTP_ACCEPT="application/json")
        assert client.get(reverse("notifications:notification_feed")).json()["alerts"] == []

    def test_feed_never_shows_someone_elses_alerts(self, client, world):
        notify(world["user_a"], title="Panic alert — ALPHA-001", alert=world["alert_a"])
        client.force_login(world["user_b"])
        assert client.get(reverse("notifications:notification_feed")).json()["alerts"] == []


def test_client_admin_role_can_view_alerts_after_seeding(db):
    call_command("seed_roles", verbosity=0)
    role = Role.objects.get(code="client-admin")
    assert role.permissions.filter(module="alert", action="view").exists()
    assert not role.permissions.filter(module="alert", action="update").exists()
