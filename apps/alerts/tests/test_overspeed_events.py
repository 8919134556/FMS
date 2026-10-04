"""OVER SPEEDING alert events (apps.alerts.overspeed): eventioval 255 ->
overspeed 1 -> one HIGH alert per overspeed episode (the device's start and end
255 records), through the real ingestion / comms-bridge paths; plus the Alert
Report, exports and bell for it, and independence from Panic and Idle."""

import datetime
import io
from decimal import Decimal

import pytest
from django.utils import timezone

from apps.accounts.models import Permission
from apps.accounts.tests.factories import UserFactory
from apps.alerts.models import Alert
from apps.alerts.overspeed import overspeed_from_eventioval
from apps.clients.tests.factories import ClientFactory
from apps.core.tests.factories import role_with
from apps.drivers.tests.factories import DriverFactory
from apps.notifications.models import Notification
from apps.tracking import comms_sync
from apps.tracking.models import TelemetryEvent
from apps.tracking.providers.base import NormalizedEvent
from apps.tracking.services import TelemetryIngestionService
from apps.tracking.tests.factories import TrackingDeviceFactory
from apps.vehicles.tests.factories import VehicleFactory

pytestmark = pytest.mark.django_db

S = datetime.timedelta(seconds=1)
M = datetime.timedelta(minutes=1)
REPORT_API = "/api/v1/alerts/report/"


def _base():
    return timezone.now() - datetime.timedelta(minutes=50)


def _vehicle(**kwargs):
    vehicle = VehicleFactory(client=kwargs.pop("client", None) or ClientFactory(), **kwargs)
    TrackingDeviceFactory(vehicle=vehicle)
    return vehicle


def _r(when, speed, *, overspeed=0, ignition=True, i=0, panic=None):
    metadata = {"location": "Siddhaiah Puranik Road, Bengaluru"}
    if overspeed is not None:
        metadata["overspeed"] = overspeed
    if panic is not None:
        metadata["panic"] = panic
    return NormalizedEvent(timestamp=when, latitude=Decimal(f"{12.98 + i * 0.0002:.6f}"), longitude=Decimal("77.53"),
                           speed=Decimal(str(speed)), ignition=ignition, metadata=metadata)


def _episode(t, speeds, *, start_speed=63, end_speed=55):
    """Like the real device: a 255 record crossing the limit, ordinary records,
    then a 255 record dropping back under it."""
    rows = [_r(t, start_speed, overspeed=1, i=0)]
    rows += [_r(t + (k + 1) * 2 * S, s, i=k + 1) for k, s in enumerate(speeds)]
    rows.append(_r(t + (len(speeds) + 1) * 2 * S, end_speed, overspeed=1, i=len(speeds) + 1))
    return rows


def _ingest(vehicle, *batches):
    for batch in batches:
        TelemetryIngestionService.ingest_events(device=vehicle.tracking_device, events=list(batch))


def _alerts(vehicle):
    return list(Alert.objects.filter(category=Alert.Category.OVER_SPEEDING, vehicle=vehicle).order_by("occurred_at"))


@pytest.fixture
def watcher():
    return UserFactory(role=role_with((Permission.Module.ALERT, Permission.Action.VIEW)))


class TestMapping:
    @pytest.mark.parametrize("value, expected", [(255, 1), ("255", 1), (0, 0), (9, 0), (239, 0), (None, None),
                                                 ("", None), ("x", None)])
    def test_eventioval_255_is_overspeed_1(self, value, expected):
        assert overspeed_from_eventioval(value) == expected

    def test_the_bridge_stores_overspeed_not_eventioval(self):
        tz = datetime.timezone(datetime.timedelta(hours=5, minutes=30))
        row = {"id": 1, "unitno": "350424067315497", "tracktime": datetime.datetime(2026, 10, 4, 12, 31, 12),
               "lat": 12.98, "lon": 77.53, "speed": 63.0, "direction": "0", "ignition": 1, "odometer": 1000,
               "gpsodometer": 0, "gpsstatus": 1, "location": "Bengaluru"}
        assert comms_sync.row_to_event({**row, "eventioval": 255}, tz).metadata["overspeed"] == 1
        assert comms_sync.row_to_event({**row, "eventioval": 9}, tz).metadata["overspeed"] == 0
        assert "overspeed" not in comms_sync.row_to_event({**row, "eventioval": None}, tz).metadata
        assert "eventioval" not in comms_sync.row_to_event({**row, "eventioval": 255}, tz).metadata


class TestEpisodes:
    def test_start_and_end_records_are_one_high_alert(self, watcher):
        v = _vehicle()
        t = _base()
        _ingest(v, _episode(t, [66, 73, 74, 66, 61]))
        [alert] = _alerts(v)
        assert alert.get_category_display() == "Over Speeding" and alert.get_severity_display() == "High"
        assert alert.occurred_at == t and alert.signal_cleared_at == t + 12 * S
        assert alert.speed == Decimal("74.00")  # the top speed of the episode
        assert Notification.objects.filter(recipient=watcher, alert=alert).count() == 1

    def test_ordinary_records_never_create_or_close(self, watcher):
        v = _vehicle()
        t = _base()
        _ingest(v, [_r(t + k * 2 * S, 80, i=k) for k in range(10)])  # fast, but no device event
        assert _alerts(v) == []
        _ingest(v, [_r(t + 30 * S, 63, overspeed=1, i=20)], [_r(t + (31 + k) * S, 70, i=21 + k) for k in range(5)])
        [alert] = _alerts(v)
        assert alert.signal_cleared_at is None  # still over the limit until the end record

    def test_redelivered_records_change_nothing(self, watcher):
        v = _vehicle()
        rows = _episode(_base(), [70, 72])
        _ingest(v, rows, rows, rows[:1])
        assert len(_alerts(v)) == 1 and Notification.objects.count() == 1

    def test_end_record_arriving_before_its_start_is_still_one_alert(self, watcher):
        """Real comms data: 13:17:48 (end) was stored before 13:17:43 (start)."""
        v = _vehicle()
        t = _base()
        rows = _episode(t, [64, 64, 64], start_speed=63, end_speed=56)
        _ingest(v, [rows[-1]], rows[1:-1], [rows[0]])
        [alert] = _alerts(v)
        assert alert.occurred_at == t and alert.signal_cleared_at == rows[-1].timestamp
        assert alert.speed == Decimal("64.00")
        assert Notification.objects.count() == 1

    def test_two_overspeeds_are_two_alerts(self, watcher):
        v = _vehicle()
        t = _base()
        _ingest(v, _episode(t, [70]), _episode(t + 5 * M, [80]))
        first, second = _alerts(v)
        assert first.signal_cleared_at < second.occurred_at
        assert (first.speed, second.speed) == (Decimal("70.00"), Decimal("80.00"))
        assert Notification.objects.count() == 2

    def test_a_lost_end_record_does_not_swallow_the_next_overspeed(self, watcher):
        v = _vehicle()
        t = _base()
        _ingest(v, [_r(t, 63, overspeed=1)], _episode(t + 40 * M, [70], start_speed=65))
        stale, new = _alerts(v)
        assert stale.signal_cleared_at == t  # closed at its last known reading, end unknown
        assert new.occurred_at == t + 40 * M and new.signal_cleared_at is not None

    def test_ignition_off_between_records_means_a_new_overspeed(self, watcher):
        v = _vehicle()
        t = _base()
        _ingest(v, [_r(t, 63, overspeed=1), _r(t + 60 * S, 0, ignition=False, i=1),
                    _r(t + 2 * M, 63, overspeed=1, i=2)])
        assert len(_alerts(v)) == 2

    def test_state_is_per_vehicle(self, watcher):
        a, b = _vehicle(), _vehicle()
        t = _base()
        _ingest(a, [_r(t, 63, overspeed=1)])
        _ingest(b, [_r(t + 3 * S, 63, overspeed=1)])  # B's record must not close A's episode
        assert len(_alerts(a)) == 1 and len(_alerts(b)) == 1
        assert _alerts(a)[0].signal_cleared_at is None

    def test_event_record_sharing_a_timestamp_with_an_ordinary_one_keeps_its_flag(self, watcher):
        """Real data: two records at 12:31:12 (eventioval 0 and 255), same position."""
        v = _vehicle()
        t = _base()
        _ingest(v, [_r(t, 60, overspeed=0), _r(t, 63, overspeed=1)])
        assert TelemetryEvent.objects.get(vehicle=v).metadata["overspeed"] == 1
        assert len(_alerts(v)) == 1


class TestIndependence:
    def test_panic_idle_and_overspeed_coexist(self, watcher):
        v = _vehicle()
        t = timezone.now() - datetime.timedelta(minutes=30)
        parked = [NormalizedEvent(timestamp=t + k * 30 * S, latitude=Decimal("12.950000"), longitude=Decimal("77.5"),
                                  speed=Decimal("0"), ignition=True,
                                  metadata={"panic": 1 if k == 2 else 0}) for k in range(13)]
        _ingest(v, parked, _episode(t + 10 * M, [75]))
        kinds = dict(Alert.objects.filter(vehicle=v).values_list("category", "severity"))
        assert kinds == {"PANIC": "CRITICAL", "IDLE": "MEDIUM", "OVER_SPEEDING": "HIGH"}


class TestNotification:
    def test_notification_names_the_alert_and_its_level(self, watcher):
        v = _vehicle(registration_number="KA01AB1234", current_driver=DriverFactory(first_name="John", last_name="Rao"))
        _ingest(v, [_r(_base(), 85, overspeed=1)])
        n = Notification.objects.get(recipient=watcher)
        assert n.title == "Over Speeding alert — KA01AB1234"
        assert "Level: High" in n.body and "85 km/h" in n.body and "Driver: John Rao" in n.body

    def test_bell_announces_it_as_high(self, client, watcher):
        _ingest(_vehicle(), [_r(_base(), 85, overspeed=1)])
        client.force_login(watcher)
        [announced] = client.get("/notifications/feed/").json()["alerts"]
        assert (announced["type"], announced["type_label"], announced["severity"], announced["level_label"]) == (
            "OVER_SPEEDING", "Over Speeding", "HIGH", "High")

    def test_old_records_do_not_notify(self, watcher):
        _ingest(_vehicle(), [_r(timezone.now() - datetime.timedelta(days=3), 85, overspeed=1)])
        assert Alert.objects.filter(category="OVER_SPEEDING").count() == 1 and not Notification.objects.exists()


class TestReportAndExports:
    def _login(self, client):
        client.force_login(UserFactory(role=role_with((Permission.Module.ALERT, Permission.Action.VIEW))))

    def test_filter_and_separate_alert_and_level(self, client):
        v = _vehicle()
        t = timezone.now() - datetime.timedelta(minutes=20)
        _ingest(v, _episode(t, [74]), [_r(t + 5 * M, 0, overspeed=0, panic=1, i=9)])
        self._login(client)
        rows = client.get(REPORT_API, {"alert_type": "OVER_SPEEDING"}).json()["alerts"]["rows"]
        assert [(r["type_label"], r["level_label"]) for r in rows] == [("Over Speeding", "High")]
        assert rows[0]["speed"] == 74.0 and rows[0]["signal_active"] is False
        assert "eventioval" not in str(rows)
        assert {r["type"] for r in client.get(REPORT_API, {"level": "HIGH"}).json()["alerts"]["rows"]} == {
            "OVER_SPEEDING"}

    def test_excel_and_pdf(self, client):
        from openpyxl import load_workbook

        _ingest(_vehicle(), _episode(timezone.now() - datetime.timedelta(minutes=20), [74]))
        self._login(client)
        xlsx = client.get("/api/v1/alerts/report/export/", {"type": "xlsx", "alert_type": "OVER_SPEEDING"})
        header, *rows = list(load_workbook(io.BytesIO(xlsx.content))["Alerts"].iter_rows(values_only=True))
        values = dict(zip(header, rows[0]))
        assert (values["Alert"], values["Level"], values["Speed (km/h)"]) == ("Over Speeding", "High", 74.0)
        assert not any("eventio" in str(h).lower() for h in header)
        pdf = client.get("/api/v1/alerts/report/export/", {"type": "pdf", "alert_type": "OVER_SPEEDING"})
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


class TestCommsBridgeEndToEnd:
    def test_comms_rows_with_eventioval_become_one_alert(self, watcher):
        v = _vehicle()
        v.tracking_device.imei = "350424067315497"
        v.tracking_device.save(update_fields=["imei"])
        now_local = timezone.now().astimezone(comms_sync._local_timezone()).replace(tzinfo=None, microsecond=0)
        t = now_local - datetime.timedelta(minutes=5)

        def row(i, sec, speed, ev, lat):
            return {"id": i, "unitno": "350424067315497", "tracktime": t + sec * S, "lat": lat, "lon": 77.53,
                    "speed": speed, "direction": "0", "ignition": 1, "odometer": 1000, "gpsodometer": 0,
                    "gpsstatus": 1, "location": "Bengaluru", "eventioval": ev}

        rows = [row(1, 0, 60.0, 0, 12.98), row(2, 0, 63.0, 255, 12.98),  # same instant, ordinary + event
                row(3, 2, 70.0, 0, 12.9802), row(4, 4, 74.0, 0, 12.9804), row(5, 6, 66.0, 9, 12.9806),
                row(6, 8, 55.0, 255, 12.9808), row(7, 10, 50.0, 0, 12.981)]
        comms_sync.apply_rows(rows, comms_sync.SyncReport())
        [alert] = _alerts(v)
        assert alert.speed == Decimal("74.00") and alert.signal_cleared_at is not None
        assert Notification.objects.filter(recipient=watcher).count() == 1
