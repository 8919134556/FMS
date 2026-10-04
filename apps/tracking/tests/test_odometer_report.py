"""Odometer Report: Device Odometer and GPS Odometer computed separately, per
vehicle per local day, from TelemetryEvent history — including the data
problems real devices produce (resets, jitter, spikes, duplicates, missing
fixes/values, gaps, midnight crossings) — plus the API, export, page and the
Live Tracking "Odometer" column."""

import datetime
import io

import pytest
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import Permission
from apps.accounts.tests.factories import UserFactory
from apps.clients.tests.factories import ClientFactory
from apps.core.tests.factories import role_with
from apps.core.utils import display_timezone
from apps.tracking import odometer_report as orp
from apps.tracking import trip_analytics
from apps.tracking.tests.factories import TelemetryEventFactory, TrackingDeviceFactory
from apps.vehicles.tests.factories import VehicleFactory

pytestmark = pytest.mark.django_db

M = datetime.timedelta(minutes=1)
S = datetime.timedelta(seconds=1)
# ~0.1 km of latitude per step: 0.0009° ≈ 0.100 km.
LAT0, LON0, STEP = 12.900000, 77.500000, 0.0009


@pytest.fixture(autouse=True)
def _no_map_tiles(settings):
    settings.TRIP_REPORT_MAP_TILE_URL = ""


def _day(offset):
    tz = display_timezone()
    return timezone.now().astimezone(tz).date() - datetime.timedelta(days=offset)


def _at(offset_days, hour, minute=0, second=0):
    tz = display_timezone()
    return datetime.datetime.combine(_day(offset_days), datetime.time(hour, minute, second), tzinfo=tz)


def _vehicle(client_obj=None):
    vehicle = VehicleFactory(client=client_obj)
    TrackingDeviceFactory(vehicle=vehicle)
    return vehicle


def _ev(vehicle, when, odometer, lat=None, lon=None, client_obj=None):
    return TelemetryEventFactory(
        device=vehicle.tracking_device, vehicle=vehicle, client=client_obj, timestamp=when, ignition=True,
        odometer=odometer, latitude=f"{lat if lat is not None else LAT0:.6f}", longitude=f"{lon if lon is not None else LON0:.6f}",
    )


def _drive(vehicle, start, *, odo_from, steps, every=M, client_obj=None, lat_from=LAT0):
    """``steps`` readings each 0.1 km further on BOTH the odometer and the GPS track."""
    for i in range(steps + 1):
        _ev(vehicle, start + i * every, round(odo_from + 0.1 * i, 1), lat=lat_from + STEP * i, client_obj=client_obj)


def _report(start_offset, end_offset=0, **kw):
    return orp.build_odometer_report(user=None, start_date=_day(start_offset), end_date=_day(end_offset), **kw)


def _row(report, vehicle, offset):
    return next(r for r in report.rows if r.vehicle.id == vehicle.id and r.date == _day(offset))


class TestDeviceOdometer:
    def test_start_end_distance_come_from_the_device_readings(self):
        v = _vehicle()
        _drive(v, _at(1, 9), odo_from=10250.0, steps=20)  # 10250.0 -> 10252.0
        row = _row(_report(1, 1), v, 1)
        assert (row.device_start, row.device_end, row.device_distance) == (10250.0, 10252.0, 2.0)
        assert row.device_start_carried is False  # no earlier reading: the day's first
        assert row.status == orp.STATUS_OK

    def test_each_day_is_separate_and_starts_where_the_previous_ended(self):
        v = _vehicle()
        _drive(v, _at(2, 9), odo_from=100.0, steps=30)  # 100.0 -> 103.0
        _drive(v, _at(1, 9), odo_from=103.5, steps=10)  # 103.5 -> 104.5 (0.5 km moved while offline)
        report = _report(2, 1)
        d2, d1 = _row(report, v, 2), _row(report, v, 1)
        assert (d2.device_start, d2.device_end, d2.device_distance) == (100.0, 103.0, 3.0)
        # Start = the last reading before the day (103.0), not the day's first (103.5).
        assert (d1.device_start, d1.device_end, d1.device_distance) == (103.0, 104.5, 1.5)
        assert d1.device_start_carried is True

    def test_travel_across_midnight_counts_on_the_day_it_is_recorded(self):
        v = _vehicle()
        _ev(v, _at(2, 23, 50), 500.0)
        _ev(v, _at(1, 0, 10), 510.0)  # 10 km driven across midnight
        _ev(v, _at(1, 0, 20), 512.0)
        report = _report(2, 1)
        assert _row(report, v, 2).device_distance == 0.0
        assert _row(report, v, 1).device_distance == 12.0
        assert _row(report, v, 2).device_distance + _row(report, v, 1).device_distance == 512.0 - 500.0

    def test_reset_adds_travel_before_and_after_and_is_flagged(self):
        """Real pattern from this fleet: 295.4 -> 0.0."""
        v = _vehicle()
        _drive(v, _at(1, 9), odo_from=293.4, steps=20)  # -> 295.4 (+2.0)
        _ev(v, _at(1, 10), 0.0)
        _drive(v, _at(1, 10, 1), odo_from=0.0, steps=23)  # -> 2.3 (+2.3)
        row = _row(_report(1, 1), v, 1)
        assert (row.device_start, row.device_end) == (293.4, 2.3)
        assert row.device_distance == pytest.approx(4.3)  # never End − Start = −291.1
        assert row.status == orp.STATUS_ISSUE
        assert any("reset" in n["text"] and n["level"] == "warning" for n in row.notes)

    def test_small_backward_step_is_not_counted_twice(self):
        v = _vehicle()
        for i, odo in enumerate([294.2, 293.7, 293.7, 294.5, 295.4]):
            _ev(v, _at(1, 9) + i * 5 * M, odo)
        row = _row(_report(1, 1), v, 1)
        assert row.device_distance == pytest.approx(1.2)  # 294.2 -> 295.4, not 1.7
        assert row.status == orp.STATUS_OK and row.notes[0]["level"] == "info"

    def test_spike_is_excluded(self):
        v = _vehicle()
        for i, odo in enumerate([100.0, 100.5, 900.0, 101.0, 101.5]):
            _ev(v, _at(1, 9) + i * M, odo)
        row = _row(_report(1, 1), v, 1)
        assert row.device_distance == pytest.approx(1.5)
        assert row.device_end == 101.5

    def test_sustained_implausible_jump_is_excluded_and_flagged(self):
        v = _vehicle()
        for i, odo in enumerate([100.0, 100.5, 600.5, 601.0]):
            _ev(v, _at(1, 9) + i * M, odo)
        row = _row(_report(1, 1), v, 1)
        assert row.device_distance == pytest.approx(1.0)  # 0.5 + 0.5, the 500 km jump excluded
        assert any("jump" in n["text"] and n["level"] == "warning" for n in row.notes)

    def test_fine_grained_counter_steps_are_not_mistaken_for_jumps(self):
        """0.1 km per second "implies" 360 km/h — normal for a 0.1-km counter."""
        v = _vehicle()
        _drive(v, _at(1, 9), odo_from=50.0, steps=30, every=S)
        row = _row(_report(1, 1), v, 1)
        assert row.device_distance == pytest.approx(3.0)
        assert row.status == orp.STATUS_OK

    def test_duplicate_readings_add_nothing(self):
        v = _vehicle()
        _drive(v, _at(1, 9), odo_from=10.0, steps=10)
        for i in range(0, 11, 2):  # duplicate every other reading (different id, same instant/value)
            TelemetryEventFactory(device=v.tracking_device, vehicle=v, timestamp=_at(1, 9) + i * M, ignition=True,
                                  odometer=round(10.0 + 0.1 * i, 1), latitude=f"{LAT0 + STEP * i + 0.000001:.6f}",
                                  longitude=f"{LON0:.6f}")
        row = _row(_report(1, 1), v, 1)
        assert row.device_distance == pytest.approx(1.0)
        assert row.gps_distance == pytest.approx(1.0, abs=0.01)

    def test_no_odometer_value_is_na_not_zero(self):
        v = _vehicle()
        for i in range(5):
            _ev(v, _at(1, 9) + i * M, None, lat=LAT0 + STEP * i)
        row = _row(_report(1, 1), v, 1)
        assert row.device_start is None and row.device_distance is None
        assert row.gps_distance == pytest.approx(0.4, abs=0.01)  # GPS still measured, independently
        assert row.status == orp.STATUS_ISSUE


class TestGpsOdometer:
    def test_gps_distance_is_computed_from_lat_lon_only(self):
        v = _vehicle()
        # The device odometer claims 50 km; the GPS track is 2.0 km: two independent values.
        for i in range(21):
            _ev(v, _at(1, 9) + i * M, 1000.0 + 2.5 * i, lat=LAT0 + STEP * i)
        row = _row(_report(1, 1), v, 1)
        assert row.device_distance == pytest.approx(50.0)
        assert row.gps_distance == pytest.approx(2.0, abs=0.01)
        assert row.difference_km == pytest.approx(48.0, abs=0.01)

    def test_gps_odometer_is_cumulative_from_the_first_fix(self):
        v = _vehicle()
        _drive(v, _at(3, 9), odo_from=0.0, steps=10)  # 1.0 km, before the report period
        _drive(v, _at(1, 9), odo_from=1.0, steps=20, lat_from=LAT0 + STEP * 10)  # +2.0 km
        row = _row(_report(1, 1), v, 1)
        assert row.gps_start == pytest.approx(1.0, abs=0.01)
        assert row.gps_end == pytest.approx(3.0, abs=0.01)
        assert row.gps_distance == pytest.approx(2.0, abs=0.01)

    def test_no_fix_points_and_jumps_are_ignored(self):
        v = _vehicle()
        _drive(v, _at(1, 9), odo_from=0.0, steps=10)
        _ev(v, _at(1, 9, 3, 30), 0.3, lat=0, lon=0)  # (0,0) no-fix
        _ev(v, _at(1, 9, 4, 30), 0.4, lat=LAT0 + 5, lon=LON0)  # a 550 km spike in 30 s
        row = _row(_report(1, 1), v, 1)
        assert row.gps_distance == pytest.approx(1.0, abs=0.01)
        assert any("GPS jump" in n["text"] for n in row.notes)

    def test_day_without_any_fix_is_na(self):
        v = _vehicle()
        for i in range(3):
            _ev(v, _at(1, 9) + i * M, 10.0 + i, lat=0, lon=0)
        row = _row(_report(1, 1), v, 1)
        assert row.gps_start is None and row.gps_distance is None
        assert row.device_distance == 2.0
        assert any("No valid GPS fix" in n["text"] for n in row.notes)

    def test_sql_gps_distance_matches_the_trip_reports_python_formula(self):
        v = _vehicle()
        coords = [(12.97, 77.59), (12.975, 77.593), (12.9801, 77.598), (12.986, 77.601), (12.99, 77.6)]
        for i, (lat, lon) in enumerate(coords):
            _ev(v, _at(1, 9) + i * 2 * M, 10.0 + i, lat=lat, lon=lon)
        expected = sum(trip_analytics.haversine_km(*coords[i], *coords[i + 1]) for i in range(len(coords) - 1))
        assert _row(_report(1, 1), v, 1).gps_distance == pytest.approx(expected, abs=0.001)


class TestReportShape:
    def test_every_vehicle_and_day_gets_a_row_with_na_for_missing_days(self):
        a, b = _vehicle(), _vehicle()
        _drive(a, _at(2, 9), odo_from=10.0, steps=10)
        _drive(b, _at(1, 9), odo_from=20.0, steps=20)
        report = _report(2, 1)
        assert len(report.rows) == 4  # 2 vehicles x 2 days
        assert _row(report, a, 1).status == orp.STATUS_NO_DATA and _row(report, a, 1).device_distance is None
        assert _row(report, b, 2).status == orp.STATUS_NO_DATA
        assert _row(report, a, 2).device_distance == 1.0 and _row(report, b, 1).device_distance == 2.0
        assert [r.date for r in report.rows] == sorted([r.date for r in report.rows], reverse=True)

    def test_vehicles_are_independent_and_filterable(self):
        a, b = _vehicle(), _vehicle()
        _drive(a, _at(1, 9), odo_from=10.0, steps=10)
        _drive(b, _at(1, 9), odo_from=5000.0, steps=30)
        only_b = _report(1, 1, vehicle_uuid=str(b.uuid))
        assert [r.vehicle.id for r in only_b.rows] == [b.id]
        assert only_b.rows[0].device_distance == 3.0
        assert only_b.summary["device_distance_km"] == 3.0

    def test_future_days_of_the_range_are_not_listed(self):
        v = _vehicle()
        _drive(v, _at(0, 0, 5), odo_from=1.0, steps=2)
        report = orp.build_odometer_report(user=None, start_date=_day(0), end_date=_day(-3))
        assert {r.date for r in report.rows} == {_day(0)}


def _view_role():
    return role_with((Permission.Module.TRACKING_DEVICE, Permission.Action.VIEW))


@pytest.mark.parametrize("range_key, days", [("today", 1), ("yesterday", 1), ("last3", 3), ("last5", 5), ("last7", 7)])
def test_api_date_filters(client, range_key, days):
    v = _vehicle()
    _drive(v, _at(1, 9), odo_from=10.0, steps=10)
    client.force_login(UserFactory(role=_view_role()))
    body = client.get(reverse("odometer-report-data"), {"range": range_key}).json()
    assert body["count"] == days
    row_dates = {r["date"] for r in body["results"]}
    assert (_day(1).isoformat() in row_dates) == (range_key != "today")
    for key in ("date", "registration_number", "device_start_km", "device_end_km", "device_distance_km",
                "gps_start_km", "gps_end_km", "gps_distance_km", "difference_km", "notes", "status"):
        assert key in body["results"][0]


def test_api_custom_range_and_scoping(client):
    mine, theirs = ClientFactory(), ClientFactory()
    own, other = _vehicle(mine), _vehicle(theirs)
    _drive(own, _at(3, 9), odo_from=1.0, steps=5, client_obj=mine)
    _drive(other, _at(3, 9), odo_from=9.0, steps=5, client_obj=theirs)
    client.force_login(UserFactory(role=_view_role(), client=mine))
    body = client.get(reverse("odometer-report-data"),
                      {"range": "custom", "from": _day(4).isoformat(), "to": _day(2).isoformat()}).json()
    assert body["range"]["start"] == _day(4).isoformat() and body["count"] == 3
    assert {r["registration_number"] for r in body["results"]} == {own.registration_number}


def test_api_requires_permission(client):
    client.force_login(UserFactory(role=None))
    assert client.get(reverse("odometer-report-data")).status_code == 403
    assert client.get(reverse("tracking:odometer_report")).status_code == 403


@pytest.mark.parametrize("file_type, magic", [("pdf", b"%PDF"), ("xlsx", b"PK")])
def test_export_downloads(client, file_type, magic):
    v = _vehicle()
    _drive(v, _at(1, 9), odo_from=10.0, steps=10)
    client.force_login(UserFactory(role=_view_role()))
    response = client.get(reverse("odometer-report-export"), {"range": "yesterday", "type": file_type})
    assert response.status_code == 200 and response.content.startswith(magic)
    assert "odometer-report_" in response["Content-Disposition"]


def test_excel_keeps_device_and_gps_in_separate_columns(client):
    from openpyxl import load_workbook

    v = _vehicle()
    _drive(v, _at(1, 9), odo_from=10.0, steps=10)
    client.force_login(UserFactory(role=_view_role()))
    content = client.get(reverse("odometer-report-export"), {"range": "yesterday", "type": "xlsx"}).content
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Summary", "Daily Odometer", "Vehicle Summary", "Daily Totals"]
    header = [c.value for c in wb["Daily Odometer"][1]]
    for column in ("Device Odometer Start (km)", "Device Odometer End (km)", "Device Odometer Distance (km)",
                   "GPS Odometer Start (km)", "GPS Odometer End (km)", "GPS Odometer Distance (km)"):
        assert column in header
    row = dict(zip(header, [c.value for c in wb["Daily Odometer"][2]]))
    assert row["Device Odometer Distance (km)"] == 1.0
    assert row["GPS Odometer Distance (km)"] == pytest.approx(1.0, abs=0.01)


def test_export_errors_are_readable(client):
    client.force_login(UserFactory(role=_view_role()))
    bad = client.get(reverse("odometer-report-export"),
                     {"range": "custom", "from": "2026-02-10", "to": "2026-02-01", "type": "pdf"})
    assert bad.status_code == 400 and "start date" in bad.json()["detail"]
    _vehicle()  # a tracked vehicle, but no history at all
    empty = client.get(reverse("odometer-report-export"), {"range": "today", "type": "xlsx"})
    assert empty.status_code == 404 and "No odometer or GPS data" in empty.json()["detail"]


def test_page_and_menu(client):
    client.force_login(UserFactory(role=_view_role()))
    response = client.get(reverse("tracking:odometer_report"))
    html = response.content.decode()
    assert response.status_code == 200
    assert response.context["active_nav_section"] == "Reports"
    assert response.context["active_nav_label"] == "Odometer Report"
    for pill in ("today", "yesterday", "last3", "last5", "last7", "custom"):
        assert f'data-range="{pill}"' in html
    assert 'id="odoVehicleFilter"' in html and 'data-report-download="xlsx"' in html
    assert "Device Odometer" in html and "GPS Odometer" in html


def test_live_tracking_column_is_the_device_odometer(client):
    client.force_login(UserFactory(role=_view_role()))
    html = client.get(reverse("tracking:live_map")).content.decode()
    assert "<th>Odometer</th>" in html and "<th>GPS Odometer</th>" not in html
