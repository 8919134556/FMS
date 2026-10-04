"""IDLE alert events (apps.alerts.idle): ignition ON + genuinely stationary,
continuously, for the vehicle's idle threshold -> one MEDIUM alert per idle
episode, through the real ingestion path; plus the Alert Report / exports /
bell for Idle, and independence from panic and from trips."""

import datetime
import io
import random
from decimal import Decimal

import pytest
from django.utils import timezone

from apps.accounts.models import Permission
from apps.accounts.tests.factories import UserFactory
from apps.alerts.models import Alert
from apps.clients.tests.factories import ClientFactory
from apps.core.tests.factories import role_with
from apps.drivers.tests.factories import DriverFactory
from apps.notifications.models import Notification
from apps.tracking import trip_report
from apps.tracking.providers.base import NormalizedEvent
from apps.tracking.services import TelemetryIngestionService
from apps.tracking.tests.factories import TrackingDeviceFactory
from apps.vehicles.tests.factories import VehicleFactory

pytestmark = pytest.mark.django_db

S = datetime.timedelta(seconds=1)
M = datetime.timedelta(minutes=1)
LAT, LON = 12.985287, 77.539518
STEP = 0.0001  # ~11 m of latitude
REPORT_API = "/api/v1/alerts/report/"


def _base():
    return timezone.now() - datetime.timedelta(hours=2)


def _vehicle(**kwargs):
    vehicle = VehicleFactory(client=kwargs.pop("client", None) or ClientFactory(), **kwargs)
    TrackingDeviceFactory(vehicle=vehicle)
    return vehicle


def _r(when, *, ignition=True, speed=0, dlat=0.0, lat=None, odometer=None, panic=None, location="MG Road, Bengaluru"):
    metadata = {"location": location}
    if panic is not None:
        metadata["panic"] = panic
    return NormalizedEvent(
        timestamp=when, latitude=Decimal(f"{lat if lat is not None else LAT + dlat:.6f}"),
        longitude=Decimal(f"{LON:.6f}"), speed=Decimal(str(speed)), ignition=ignition,
        odometer=None if odometer is None else Decimal(str(odometer)), metadata=metadata,
    )


def _ingest(vehicle, readings, batch=10):
    """Like the comms bridge: several small batches, in order."""
    for i in range(0, len(readings), batch):
        TelemetryIngestionService.ingest_events(device=vehicle.tracking_device, events=readings[i:i + batch])


def _parked(start, minutes, *, every=30 * S, jitter=True, dlat=0.0, **kwargs):
    """Ignition ON, stationary at ``dlat``: a reading every ``every`` for
    ``minutes`` (inclusive), with a few metres of GPS drift."""
    count = int(minutes * 60 / every.total_seconds())
    rng = random.Random(7)
    return [_r(start + i * every, dlat=dlat + (rng.uniform(-0.00003, 0.00003) if jitter else 0.0), **kwargs)
            for i in range(count + 1)]


def _drive(start, seconds, *, from_dlat=0.0, every=10 * S, speed=30):
    """Moving ~80 m per 10 s."""
    return [_r(start + i * every, speed=speed, dlat=from_dlat + (i + 1) * 7 * STEP)
            for i in range(int(seconds / every.total_seconds()))]


def _idles(vehicle):
    return list(Alert.objects.filter(category=Alert.Category.IDLE, vehicle=vehicle).order_by("occurred_at"))


@pytest.fixture
def watcher():
    return UserFactory(role=role_with((Permission.Module.ALERT, Permission.Action.VIEW)))


class TestUserScenarios:
    def test_1_stationary_3_minutes_is_no_alert(self, watcher):
        v = _vehicle()
        _ingest(v, _parked(_base(), 3))
        assert _idles(v) == []

    def test_2_stationary_exactly_5_minutes_is_one_medium_alert(self, watcher):
        v = _vehicle()
        t = _base()
        _ingest(v, _parked(t, 5))
        [alert] = _idles(v)
        assert alert.severity == Alert.Severity.MEDIUM and alert.status == Alert.Status.OPEN
        assert alert.occurred_at == t and alert.triggered_at == t + 5 * M
        assert alert.signal_cleared_at is None  # still idling
        assert Notification.objects.filter(recipient=watcher, alert=alert).count() == 1

    def test_3_stationary_10_minutes_is_still_one_alert(self, watcher):
        v = _vehicle()
        t = _base()
        _ingest(v, _parked(t, 10))
        [alert] = _idles(v)
        assert alert.last_signal_at == t + 10 * M
        assert Notification.objects.count() == 1  # one notification -> one sound

    def test_4_stationary_4_minutes_then_moves_resets_the_timer(self, watcher):
        v = _vehicle()
        t = _base()
        readings = _parked(t, 4) + _drive(t + 4 * M + 10 * S, 60)
        # stops again: a NEW timer — 4 more minutes is still not enough
        stop = readings[-1].timestamp + 10 * S
        readings += _parked(stop, 4, jitter=False, dlat=float(readings[-1].latitude) - LAT)
        _ingest(v, readings)
        assert _idles(v) == []

    def test_5_idle_then_moves_ends_the_event(self, watcher):
        v = _vehicle()
        t = _base()
        parked = _parked(t, 12, every=30 * S)
        departure = t + 12 * M + 30 * S
        _ingest(v, parked + _drive(departure - 10 * S, 90))
        [alert] = _idles(v)
        assert alert.triggered_at == t + 5 * M
        assert t + 12 * M <= alert.signal_cleared_at <= departure
        duration = alert.signal_cleared_at - alert.occurred_at
        assert duration >= 12 * M  # the real idle time, not the 5-minute threshold

    def test_6_ignition_off_before_5_minutes_is_no_alert(self, watcher):
        v = _vehicle()
        t = _base()
        readings = _parked(t, 3) + [_r(t + 3 * M + 30 * S, ignition=False)]
        readings += [_r(t + 4 * M + i * 30 * S, ignition=False) for i in range(8)]
        _ingest(v, readings)
        assert _idles(v) == []

    def test_7_a_few_metres_of_gps_drift_is_still_stationary(self, watcher):
        v = _vehicle()
        t = _base()
        drift = [0, 0.00003, -0.00002, 0.00004, -0.00003]  # up to ~4 m
        _ingest(v, [_r(t + i * 30 * S, dlat=drift[i % 5]) for i in range(13)])
        assert len(_idles(v)) == 1

    def test_8_one_speed_spike_is_not_movement(self, watcher):
        v = _vehicle()
        t = _base()
        speeds = [0, 0, 2, 0, 0, 0, 18, 0, 0, 0, 0, 0, 0]  # a 2 km/h blip and even an 18 km/h spike
        _ingest(v, [_r(t + i * 30 * S, speed=s) for i, s in enumerate(speeds)])
        [alert] = _idles(v)
        assert alert.occurred_at == t

    def test_9_states_are_per_vehicle(self, watcher):
        idle_v, moving_v = _vehicle(), _vehicle()
        t = _base()
        _ingest(idle_v, _parked(t, 6))
        _ingest(moving_v, _drive(t, 6 * 60))
        assert len(_idles(idle_v)) == 1 and _idles(moving_v) == []

    def test_10_panic_and_idle_are_independent(self, watcher):
        v = _vehicle()
        t = _base()
        readings = _parked(t, 6)
        readings[3] = _r(readings[3].timestamp, panic=1)
        readings[4] = _r(readings[4].timestamp, panic=0)
        _ingest(v, readings)
        panic = Alert.objects.get(category=Alert.Category.PANIC, vehicle=v)
        [idle] = _idles(v)
        assert panic.severity == Alert.Severity.CRITICAL and idle.severity == Alert.Severity.MEDIUM
        assert panic.pk != idle.pk
        assert Notification.objects.filter(alert=panic).count() == 1
        assert Notification.objects.filter(alert=idle).count() == 1


class TestRobustness:
    def test_a_gps_jump_is_not_movement(self, watcher):
        v = _vehicle()
        t = _base()
        readings = _parked(t, 6, jitter=False)
        readings[5] = _r(readings[5].timestamp, dlat=0.004)  # one fix ~440 m away, then back
        _ingest(v, readings)
        [alert] = _idles(v)
        assert alert.occurred_at == t

    def test_missing_data_is_never_counted_as_idling(self, watcher):
        v = _vehicle()
        t = _base()
        readings = _parked(t, 3) + _parked(t + 3 * M + 10 * M, 3)  # 10-minute hole in the middle
        _ingest(v, readings)
        assert _idles(v) == []

    def test_odometer_creep_while_parked_does_not_end_the_idle(self, watcher):
        """Real device: +0.2 km over 24 min within 23 m with the engine running."""
        v = _vehicle()
        t = _base()
        readings = [_r(t + i * 30 * S, odometer=70.2 if i < 10 else 70.4) for i in range(25)]
        _ingest(v, readings)
        [alert] = _idles(v)
        assert alert.signal_cleared_at is None

    def test_odometer_counter_jump_does_not_end_the_idle(self, watcher):
        v = _vehicle()
        t = _base()
        _ingest(v, [_r(t + i * 30 * S, odometer=70.4 if i < 6 else 108605.0) for i in range(14)])
        [alert] = _idles(v)
        assert alert.signal_cleared_at is None

    def test_without_a_gps_fix_the_odometer_shows_movement(self, watcher):
        v = _vehicle()
        t = _base()
        readings = [_r(t + i * 30 * S, odometer=10.0) for i in range(4)]
        readings += [_r(t + (4 + i) * 30 * S, lat=0, odometer=10.0 + 0.3 * (i + 1)) for i in range(10)]
        _ingest(v, [NormalizedEvent(**{**r.__dict__, "longitude": Decimal(0)}) if r.latitude == 0 else r
                    for r in readings])
        assert _idles(v) == []

    def test_out_of_order_and_duplicate_batches_give_the_same_single_alert(self, watcher):
        v = _vehicle()
        t = _base()
        readings = _parked(t, 9)
        batches = [readings[i:i + 4] for i in range(0, len(readings), 4)]
        order = [1, 0, 3, 2, 4, 1]  # shuffled, and batch 1 delivered twice
        for i in order:
            TelemetryIngestionService.ingest_events(device=v.tracking_device, events=batches[i])
        [alert] = _idles(v)
        assert alert.occurred_at == t and alert.last_signal_at == t + 9 * M
        assert Notification.objects.count() == 1

    def test_two_idles_separated_by_a_drive_are_two_alerts(self, watcher):
        v = _vehicle()
        t = _base()
        first = _parked(t, 6, jitter=False)
        drive = _drive(t + 6 * M + 10 * S, 120)
        here = float(drive[-1].latitude) - LAT
        second = _parked(drive[-1].timestamp + 10 * S, 6, jitter=False, dlat=here)
        _ingest(v, first + drive + second)
        a, b = _idles(v)
        assert a.signal_cleared_at <= b.occurred_at
        assert Notification.objects.count() == 2

    def test_old_history_does_not_notify(self, watcher):
        v = _vehicle()
        _ingest(v, _parked(timezone.now() - datetime.timedelta(days=3), 6))
        assert len(_idles(v)) == 1 and not Notification.objects.exists()


class TestConfiguration:
    def test_idle_threshold_is_per_vehicle(self, watcher):
        quick, slow = _vehicle(idle_alert_minutes=5), _vehicle(idle_alert_minutes=10)
        t = _base()
        for v in (quick, slow):
            _ingest(v, _parked(t, 8))
        assert len(_idles(quick)) == 1 and _idles(slow) == []
        _ingest(slow, _parked(t + 8 * M + 30 * S, 2))
        [alert] = _idles(slow)
        assert alert.triggered_at == t + 10 * M

    def test_stationary_radius_is_the_vehicles_movement_distance(self, watcher):
        tight, loose = _vehicle(trip_min_distance_m=20), _vehicle(trip_min_distance_m=200)
        t = _base()
        creep = [_r(t + i * 30 * S, dlat=i * 0.00006) for i in range(13)]  # ~6.7 m per reading, 80 m total
        for v in (tight, loose):
            _ingest(v, creep)
        assert _idles(tight) == [] and len(_idles(loose)) == 1

    def test_default_threshold(self):
        assert VehicleFactory().idle_alert_minutes == 5


class TestNotification:
    def test_notification_is_medium_and_says_what_happened(self, watcher):
        v = _vehicle(registration_number="KA01AB1234", current_driver=DriverFactory(first_name="John", last_name="Rao"))
        _ingest(v, _parked(_base(), 6))
        n = Notification.objects.get(recipient=watcher)
        assert n.title == "Idle alert — KA01AB1234" and n.level == Notification.Level.WARNING
        assert "Stationary with ignition ON for 5 min or more" in n.body and "Driver: John Rao" in n.body

    def test_bell_announces_it_with_its_severity(self, client, watcher):
        v = _vehicle()
        _ingest(v, _parked(_base(), 6))
        client.force_login(watcher)
        [announced] = client.get("/notifications/feed/").json()["alerts"]
        assert (announced["type"], announced["severity"]) == ("IDLE", "MEDIUM")
        assert announced["message"].startswith("Stationary with ignition ON")

    def test_other_clients_users_are_not_notified(self):
        a, b = ClientFactory(), ClientFactory()
        view = role_with((Permission.Module.ALERT, Permission.Action.VIEW))
        user_a, user_b = UserFactory(role=view, client=a), UserFactory(role=view, client=b)
        _ingest(_vehicle(client=a), _parked(_base(), 6))
        assert set(Notification.objects.values_list("recipient_id", flat=True)) == {user_a.pk}
        assert user_b.pk not in set(Notification.objects.values_list("recipient_id", flat=True))


class TestIndependentFromTrips:
    def test_idling_mid_trip_does_not_end_or_split_the_trip(self, watcher):
        v = _vehicle()
        t = _base()
        readings = _drive(t, 120)
        here = float(readings[-1].latitude) - LAT
        stop = readings[-1].timestamp + 10 * S
        readings += _parked(stop, 7, jitter=False, dlat=here)
        readings += _drive(stop + 7 * M + 10 * S, 120, from_dlat=here)
        readings.append(_r(readings[-1].timestamp + 10 * S, ignition=False, dlat=float(readings[-1].latitude) - LAT))
        _ingest(v, readings)
        assert len(_idles(v)) == 1
        trips, _ = trip_report.compute_vehicle_trips(user=None, start_date=t.date(), end_date=timezone.now().date(),
                                                     vehicle_uuid=str(v.uuid))
        assert len(trips) == 1


class TestReportAndExports:
    def _staff(self, client):
        client.force_login(UserFactory(role=role_with((Permission.Module.ALERT, Permission.Action.VIEW))))

    def _data(self):
        v = _vehicle(registration_number="KA01AB1234")
        t = timezone.now() - datetime.timedelta(minutes=40)
        readings = _parked(t, 12, jitter=False)
        readings[2] = _r(readings[2].timestamp, panic=1)
        readings[3] = _r(readings[3].timestamp, panic=0)
        _ingest(v, readings + _drive(t + 12 * M + 10 * S, 60))
        return v

    def test_alert_and_level_filters(self, client):
        self._data()
        self._staff(client)
        body = client.get(REPORT_API).json()
        by_type = body["summary"]["by_type"]
        assert (by_type["PANIC"], by_type["IDLE"]) == (1, 1) and sum(by_type.values()) == 2
        assert body["summary"]["by_level"]["MEDIUM"] == 1 and body["summary"]["by_level"]["CRITICAL"] == 1
        idle_rows = client.get(REPORT_API, {"alert_type": "IDLE"}).json()["alerts"]["rows"]
        assert [r["type"] for r in idle_rows] == ["IDLE"]
        medium = client.get(REPORT_API, {"level": "MEDIUM"}).json()["alerts"]["rows"]
        assert [(r["type_label"], r["level_label"]) for r in medium] == [("Idle", "Medium")]  # two separate values
        assert client.get(REPORT_API, {"level": "EXTREME"}).status_code == 400
        row = idle_rows[0]
        assert row["duration_text"].startswith("12m") and row["triggered_at"] != row["occurred_at"]
        assert row["signal_active"] is False

    def test_excel_has_the_idle_lifecycle_columns(self, client):
        from openpyxl import load_workbook

        self._data()
        self._staff(client)
        response = client.get("/api/v1/alerts/report/export/", {"type": "xlsx", "alert_type": "IDLE"})
        header, *rows = list(load_workbook(io.BytesIO(response.content))["Alerts"].iter_rows(values_only=True))
        for column in ("Alert", "Level", "Start time", "Alert time", "End time", "Duration", "Duration (s)"):
            assert column in header
        [row] = rows
        values = dict(zip(header, row))
        assert values["Alert"] == "Idle" and values["Level"] == "Medium"
        assert values["Duration (s)"] >= 12 * 60 and values["End time"] is not None

    def test_pdf_lists_idle_alerts(self, client):
        self._data()
        self._staff(client)
        response = client.get("/api/v1/alerts/report/export/", {"type": "pdf", "level": "MEDIUM"})
        assert response.status_code == 200 and response.content.startswith(b"%PDF")
