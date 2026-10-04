"""PDF / Excel for the Location Data Report — EVERY matching GPS record
(streamed from apps.tracking.location_report.iter_records, never just the
page on screen), with the report's analysis and data-quality summary, built on
the Trip Report's shared blocks (apps.tracking.trip_exports).

Only columns that hold data in the selection are included (the others exist
in the schema but are empty for this device/provider)."""

from django.utils import timezone

from apps.tracking import location_report as lr
from apps.tracking import trip_analytics as ta
from apps.tracking.trip_exports import (
    DASH,
    FMT_COORD,
    FMT_DATE,
    FMT_DT,
    FMT_INT,
    FMT_TIME,
    _Pdf,
    _summary_sheet,
    _xl_dt,
    _Xlsx,
    fmt_dt,
    fmt_duration,
    fmt_km,
    fmt_speed,
)

FLAG_SHORT = {
    "no_fix": "No fix", "device_no_fix": "Device no-fix", "invalid_coords": "Invalid coords",
    "duplicate": "Duplicate", "gap": "Gap", "jump": "GPS jump", "invalid_speed": "Invalid speed",
}
QUALITY_SCOPE = {"": "All records", "valid": "Valid records only", "issues": "Flagged records only"}


def _meta(selection, quality_filter):
    s, e = selection.start_date, selection.end_date
    span = f"{s:%d %b %Y}" if s == e else f"{s:%d %b %Y} – {e:%d %b %Y}"
    vehicle = "All vehicles" if selection.all_vehicles else selection.vehicles[0].registration_number
    return ta._meta(
        title="Location Data Report", subtitle=f"{vehicle} · {span}", period_start=selection.period_start,
        period_end=selection.period_end, period_label=f"{ta.RANGE_LABELS.get(selection.range_key, '')} ({span})",
        filters_label=f"Vehicle: {vehicle} · {QUALITY_SCOPE.get(quality_filter, 'All records')}",
        now=timezone.now(), tz=selection.tz, vehicles=selection.vehicles,
    )


def _analysis_rows(analysis, quality, tz):
    """[(label, value)] — only metrics that exist for this selection."""
    q = quality
    rows = [
        ("Total GPS records", f"{q['records']:,}"),
        ("First record", fmt_dt(q["first_ts"], tz)),
        ("Last record", fmt_dt(q["last_ts"], tz)),
    ]
    if analysis:
        a = analysis
        rows += [
            ("Tracking duration", fmt_duration(a["tracking_seconds"])),
            ("GPS distance travelled", fmt_km(a["gps_distance_km"], 2)),
            ("Moving duration", fmt_duration(a["moving_seconds"])),
            ("Idle duration (ignition on, not moving)", fmt_duration(a["idle_seconds"])),
            ("Stopped duration (idle + ignition off)", fmt_duration(a["stopped_seconds"])),
            ("Ignition ON duration", fmt_duration(a["ignition_on_seconds"])),
            ("Ignition OFF duration", fmt_duration(a["ignition_off_seconds"])),
            ("Ignition ON / OFF events", f"{a['ignition_on_events']} / {a['ignition_off_events']}"),
            ("Maximum speed", fmt_speed(a["max_speed_kmh"])
             + (f" at {fmt_dt(a['max_speed_at'], tz, '%d %b %H:%M:%S')}" if a["max_speed_at"] else "")),
            ("Average moving speed", fmt_speed(a["avg_moving_speed_kmh"])),
            ("Stops (≥ 1 min)", f"{a['stops']}"
             + (f", longest {fmt_duration(a['longest_stop_seconds'])}" if a["longest_stop_seconds"] else "")),
        ]
        if a["no_data_seconds"]:
            rows.append((f"Time without data (gaps > {a['gap_threshold_minutes']} min)", fmt_duration(a["no_data_seconds"])))
    return rows


def _quality_rows(quality):
    q = quality
    gap = q["gap"]
    return [
        ("Valid GPS records", f"{q['valid_fixes']:,}"),
        ("No GPS fix (0,0 coordinates)", f"{q['no_fix']:,}"),
        ("Device reports no GPS fix", f"{q['device_no_fix']:,}"),
        ("Invalid coordinates", f"{q['invalid_coords']:,}"),
        ("Duplicate timestamps", f"{q['duplicate']:,}"),
        ("GPS data gaps (> 15 min)", f"{gap:,}"
         + (f", longest {fmt_duration(q['longest_gap_seconds'])}" if gap else "")),
        ("GPS jumps (impossible movement)", f"{q['jump']:,}"),
        ("Invalid speed values", f"{q['invalid_speed']:,}"),
        ("Records with any flag", f"{q['flagged']:,}"),
    ]


def _flags_short(row):
    out = []
    for key, label in FLAG_SHORT.items():
        if row.get(key):
            out.append(f"{label} {int(row['gap_seconds'] // 60)}m" if key == "gap" else label)
    return ", ".join(out) or "Valid"


def _num(value, places):
    return f"{float(value):.{places}f}" if value is not None else DASH


def location_report_pdf(selection, quality, analysis, *, quality_filter=""):
    from reportlab.platypus import PageBreak, Spacer

    total = quality["records"] if not quality_filter else (
        quality["flagged"] if quality_filter == "issues" else quality["records"] - quality["flagged"])
    if total > lr.PDF_MAX_ROWS:
        raise ta.ReportError(
            f"This selection has {total:,} GPS records — too many for a readable PDF (maximum {lr.PDF_MAX_ROWS:,}). "
            "Download Excel for the complete data, or choose a shorter period.", status=413)
    meta = _meta(selection, quality_filter)
    tz = selection.tz
    pdf = _Pdf(meta)
    story = pdf.title_block()

    a = analysis or {}
    cards = [
        ("GPS records", f"{quality['records']:,}", f"{quality['valid_fixes']:,} valid"),
        ("GPS distance", fmt_km(a.get("gps_distance_km"), 2), "from latitude/longitude"),
        ("Tracking duration", fmt_duration(a.get("tracking_seconds")), "first → last record"),
        ("Maximum speed", fmt_speed(a.get("max_speed_kmh")), f"avg {fmt_speed(a.get('avg_moving_speed_kmh'))} moving"),
        ("Moving", fmt_duration(a.get("moving_seconds")), ""),
        ("Stopped", fmt_duration(a.get("stopped_seconds")), "idle + ignition off"),
        ("Ignition ON", fmt_duration(a.get("ignition_on_seconds")),
         f"{a.get('ignition_on_events', 0)} ON / {a.get('ignition_off_events', 0)} OFF events"),
        ("Flagged records", f"{quality['flagged']:,}", "see Data Quality"),
    ]
    story += [pdf.section("1. Analysis Summary"), pdf.kpi_grid(cards, cols=4),
              pdf.key_value(_analysis_rows(analysis, quality, tz))]
    story += [pdf.section("2. Data Quality", "Records are never removed: each problem record is listed with its flag."),
              pdf.key_value(_quality_rows(quality))]

    show_vehicle = selection.all_vehicles
    headers = ["#", "Date / time"] + (["Vehicle"] if show_vehicle else []) + [
        "Latitude", "Longitude", "Speed km/h", "Ignition", "GPS", "Odometer km", "Location", "Status"]
    widths = [0.45, 1.25] + ([0.9] if show_vehicle else []) + [0.75, 0.75, 0.6, 0.55, 0.45, 0.7, 2.6, 1.0]
    rows = []
    for i, r in enumerate(lr.iter_records(selection, quality=quality_filter), start=1):
        loc = r["location"] or ""
        rows.append([i, fmt_dt(r["ts"], tz, "%d %b %Y %H:%M:%S")] + ([r["registration_number"]] if show_vehicle else []) + [
            _num(r["lat"], 6), _num(r["lon"], 6), _num(r["speed"], 1),
            {True: "On", False: "Off"}.get(r["ignition"], DASH),
            {"1": "Fix", "0": "No fix"}.get(r["gps_status"], DASH),
            _num(r["odometer"], 1), (loc[:57] + "…") if len(loc) > 58 else (loc or DASH), _flags_short(r),
        ])
    story += [PageBreak(), pdf.section("3. Location Records",
                                       f"{len(rows):,} records, oldest first ({QUALITY_SCOPE.get(quality_filter)}).")]
    if rows:
        numeric = {"#", "Latitude", "Longitude", "Speed km/h", "Odometer km"}
        story.append(pdf.table(headers, rows, widths, dense=True,
                               align_right=[i for i, h in enumerate(headers) if h in numeric]))
    else:
        story.append(pdf.p("No records match this selection."))
    story += [Spacer(1, 10), pdf.p(
        "Distance and durations use the Trip Report's analysis rules: GPS distance between consecutive valid fixes "
        f"(hops faster than {ta.MAX_PLAUSIBLE_SPEED_KMH} km/h ignored); Moving at ≥ "
        f"{a.get('moving_threshold_kmh', 5)} km/h; gaps of more than {a.get('gap_threshold_minutes', 15)} min count as "
        "“no data”. GPS = the device's own fix status.", "note")]
    return pdf.build(story)


# Excel ----------------------------------------------------------------------

_XL_COLUMNS = [
    # key, header, number format, width, optional-column key (None = always)
    ("n", "#", FMT_INT, 8, None),
    ("ts", "GPS timestamp", FMT_DT, 20, None),
    ("date", "Date", FMT_DATE, 12, None),
    ("time", "Time", FMT_TIME, 10, None),
    ("registration_number", "Vehicle", None, 14, None),
    ("lat", "Latitude", FMT_COORD, 12, None),
    ("lon", "Longitude", FMT_COORD, 12, None),
    ("speed", "Speed (km/h)", "0.00", 11, None),
    ("heading", "Heading (°)", FMT_INT, 10, "heading"),
    ("altitude", "Altitude (m)", "0.0", 11, "altitude"),
    ("ignition", "Ignition", None, 9, None),
    ("gps_status", "GPS status", None, 10, "gps_status"),
    ("satellite_count", "Satellites", FMT_INT, 10, "satellite_count"),
    ("odometer", "Odometer (km)", "#,##0.0", 13, None),
    ("gps_odometer_km", "Device trip meter (km)", "0.000", 14, "gps_odometer_km"),
    ("engine_hours", "Engine hours", "0.0", 12, "engine_hours"),
    ("battery_voltage", "Battery (V)", "0.00", 11, "battery_voltage"),
    ("external_power", "External power", None, 12, "external_power"),
    ("signal_strength", "Signal", FMT_INT, 9, "signal_strength"),
    ("location", "Location / address", None, 70, "location"),
    ("gap_minutes", "Minutes since previous record", "0.0", 14, None),
    ("status", "Data status", None, 30, None),
]


def _xl_value(key, r, tz, n):
    if key == "n":
        return n
    if key in ("ts", "date", "time"):
        local = _xl_dt(r["ts"], tz)
        return local if key == "ts" else (local.date() if key == "date" else local.time())
    if key == "ignition":
        return {True: "On", False: "Off"}.get(r["ignition"])
    if key == "external_power":
        return {True: "Yes", False: "No"}.get(r["external_power"])
    if key == "gps_status":
        return {"1": "Fix", "0": "No fix"}.get(r["gps_status"], r["gps_status"])
    if key == "gps_odometer_km":
        return float(r["gps_odometer_km"]) if r["gps_odometer_km"] not in (None, "") else None
    if key == "gap_minutes":
        return round(float(r["gap_seconds"]) / 60, 1) if r["gap_seconds"] is not None else None
    if key == "status":
        return "; ".join(lr.record_flags(r)) or "Valid"
    if key == "location":
        return r["location"] or None
    value = r.get(key)
    return float(value) if isinstance(value, float) else value


def _stream_sheet(x, title, columns, records, tz):
    """A typed, filterable sheet written row by row (write-only workbook)."""
    from openpyxl.utils import get_column_letter

    ws = x.wb.create_sheet(title)
    for j, (_k, _h, _f, width, _o) in enumerate(columns):
        ws.column_dimensions[get_column_letter(j + 1)].width = width
    ws.freeze_panes = "C2"
    ws.append([x.cell(ws, header, font=x.header_font, fill=x.header_fill,
                      align=x.Alignment(vertical="center", wrap_text=True)) for _k, header, _f, _w, _o in columns])
    count = 0
    for count, r in enumerate(records, start=1):
        if count > lr.EXCEL_MAX_ROWS:
            raise ta.ReportError(
                f"This selection has more than {lr.EXCEL_MAX_ROWS:,} GPS records — beyond what one Excel sheet can "
                "hold. Please choose a shorter period or a single vehicle.", status=413)
        ws.append([x.cell(ws, v, fmt=fmt) if v is not None else None
                   for (k, _h, fmt, _w, _o), v in ((c, _xl_value(c[0], r, tz, count)) for c in columns)])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(count, 1) + 1}"
    return count


def location_report_xlsx(selection, quality, analysis, *, quality_filter=""):
    meta = _meta(selection, quality_filter)
    x = _Xlsx(meta)
    tz = selection.tz
    present = {c for c in lr.OPTIONAL_COLUMNS if quality.get(c)}
    columns = [c for c in _XL_COLUMNS if c[4] is None or c[4] in present]

    metrics = [(label, value, None) for label, value in _analysis_rows(analysis, quality, tz)]
    metrics += [(label, value, None) for label, value in _quality_rows(quality)]
    notes = [
        f"Location Data sheet: every record of the selection ({QUALITY_SCOPE.get(quality_filter)}), oldest first, "
        "exactly as stored in the GPS history. Columns without any data for this selection are omitted.",
        "Data status flags problem records instead of removing them. Distances/durations follow the Trip Report's "
        f"analysis rules (GPS jumps faster than {ta.MAX_PLAUSIBLE_SPEED_KMH} km/h ignored; gaps > "
        f"{int(ta.MAX_SAMPLE_GAP.total_seconds() // 60)} min = no data).",
    ]
    _summary_sheet(x, "Summary", metrics, [], notes)
    _stream_sheet(x, "Location Data", columns, lr.iter_records(selection, quality=quality_filter), tz)
    if quality["flagged"] and quality_filter != "valid":
        _stream_sheet(x, "Data Quality", columns, lr.iter_records(selection, quality="issues"), tz)
    return x.save()
