"""Trip Report downloads (PDF/Excel) + the MAP popup's per-trip analysis:
the metrics are computed from real readings (never invented), downloads
cover exactly the requested period/trip, bad input is a readable 4xx, and
client scoping holds for every new endpoint."""

import datetime
import io
import zipfile
from decimal import Decimal

import pytest
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import Permission
from apps.accounts.tests.factories import UserFactory
from apps.audit.models import AuditLog
from apps.clients.tests.factories import ClientFactory
from apps.core.tests.factories import role_with
from apps.core.utils import display_timezone
from apps.drivers.tests.factories import DriverFactory
from apps.tracking import trip_analytics
from apps.tracking.tests.factories import TelemetryEventFactory, TrackingDeviceFactory
from apps.vehicles.models import VehicleDriverAssignment
from apps.vehicles.tests.factories import VehicleFactory

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def _no_map_tiles(settings):
    """Never reach out to a tile server from the test suite — the route is
    still drawn (on a plain grid), which is what these tests care about."""
    settings.TRIP_REPORT_MAP_TILE_URL = ""


def _view_role():
    return role_with((Permission.Module.TRACKING_DEVICE, Permission.Action.VIEW))


def _yesterday_at(hour, minute=0):
    tz = display_timezone()
    day = timezone.now().astimezone(tz).date() - datetime.timedelta(days=1)
    return datetime.datetime.combine(day, datetime.time(hour, minute), tzinfo=tz)


# (minute offset, ignition, speed km/h, odometer km, latitude, longitude)
# Ignition on at minute 0 (parked, no GPS fix yet), departs at minute 1: drive
# 2 min, stand 2 min (engine on), engine off 2 min (inside the 5-min grace),
# drive 3 min, ignition off. The TRIP is minute 1 -> 10 — ignition ON is only
# a candidate; the trip starts where movement is confirmed to have begun.
READINGS = [
    (0, True, 0, "100.0", "0", "0"),  # no GPS fix yet (0,0) — must not be mapped
    (1, True, 30, "100.0", "12.970000", "77.590000"),
    (2, True, 40, "101.0", "12.975000", "77.590000"),
    (3, True, 0, "101.5", "12.978000", "77.590000"),
    (4, True, 0, "101.5", "12.978000", "77.590000"),
    (5, False, 0, "101.5", "12.978000", "77.590000"),
    (7, True, 50, "102.0", "12.982000", "77.590000"),
    (9, True, 10, "103.0", "12.990000", "77.590000"),
    (10, False, 0, "103.5", "12.994000", "77.590000"),
]


def _seed_trip(client_obj=None, start=None):
    vehicle = VehicleFactory(client=client_obj)
    TrackingDeviceFactory(vehicle=vehicle)
    start = start or _yesterday_at(10)
    for minute, ignition, speed, odometer, lat, lon in READINGS:
        TelemetryEventFactory(
            device=vehicle.tracking_device, vehicle=vehicle, client=client_obj,
            timestamp=start + datetime.timedelta(minutes=minute), ignition=ignition, speed=speed,
            odometer=odometer, latitude=lat, longitude=lon, metadata={"location": f"Point {minute}"},
        )
    return vehicle, start + datetime.timedelta(minutes=1)  # the trip starts at the departure


def _pdf_section_titles(monkeypatch):
    """Records every section heading a PDF build emits (reportlab compresses
    page streams and embeds glyph subsets, so the text isn't greppable)."""
    from apps.tracking import trip_exports

    titles = []
    original = trip_exports._Pdf.section

    def recording(self, title, note=""):
        titles.append(title)
        return original(self, title, note)

    monkeypatch.setattr(trip_exports._Pdf, "section", recording)
    return titles


class TestTripAnalysisMetrics:
    def test_every_metric_is_derived_from_the_readings(self):
        vehicle, start = _seed_trip()
        user = UserFactory(role=_view_role())

        report = trip_analytics.build_trip_report(user=user, vehicle_uuid=str(vehicle.uuid), start_at=start)
        a = report.record.analysis

        assert report.record.trip.status == "COMPLETED"
        assert a.duration_seconds == 540  # departure (minute 1) -> ignition off (minute 10)
        assert a.point_count == 8 and a.fix_count == 8  # the parked no-fix minute is before the trip
        # moving: 1->2, 2->3, 7->9, 9->10 ; idle: 3->4, 4->5 ; engine off: 5->7
        assert a.moving_seconds == 300
        assert a.idle_seconds == 120
        assert a.engine_off_seconds == 120
        assert a.no_data_seconds == 0
        assert a.distance_km == 3.5 and a.distance_source == "Odometer"  # same figure the page shows
        assert a.max_speed_kmh == 50.0
        assert a.avg_speed_kmh == pytest.approx(3.5 / (540 / 3600))
        assert a.avg_moving_speed_kmh == pytest.approx(42.0)
        # One stop: 3->7 (idle then engine off). Parked time before departure isn't a stop of the trip.
        assert [s.duration_seconds for s in a.stops] == [240]
        assert a.stops[0].engine_off is True
        assert a.gps_distance_km > 0
        assert len(report.points) == 8

    def test_trip_id_is_stable_and_readable(self):
        vehicle, start = _seed_trip()
        user = UserFactory(role=_view_role())
        report = trip_analytics.build_trip_report(user=user, vehicle_uuid=str(vehicle.uuid), start_at=start)
        assert report.record.trip_id == f"{vehicle.registration_number}-{start:%Y%m%d-%H%M%S}"

    def test_implausible_moving_speed_is_withheld_not_published(self):
        """Odometer moved while every speed reading was below the moving
        threshold -> distance / tiny moving time would exceed the top speed."""
        vehicle = VehicleFactory()
        TrackingDeviceFactory(vehicle=vehicle)
        start = _yesterday_at(9)
        for minute, ignition, speed, odo in [(0, True, 0, "10"), (1, True, 6, "10"), (2, True, 1, "12"),
                                             (30, True, 0, "12"), (31, False, 0, "12")]:
            TelemetryEventFactory(device=vehicle.tracking_device, vehicle=vehicle, ignition=ignition, speed=speed,
                                  odometer=odo, timestamp=start + datetime.timedelta(minutes=minute))
        # Confirmed by the odometer (10 -> 12 km); the trip starts at minute 1,
        # the last reading before the odometer advanced.
        report = trip_analytics.build_trip_report(
            user=UserFactory(role=_view_role()), vehicle_uuid=str(vehicle.uuid),
            start_at=start + datetime.timedelta(minutes=1),
        )
        assert report.record.analysis.avg_moving_speed_kmh is None
        assert report.record.analysis.avg_speed_kmh is not None

    def test_driver_comes_from_the_assignment_in_force_on_the_trip_date(self):
        vehicle, start = _seed_trip()
        then_driver, now_driver = DriverFactory(), DriverFactory()
        vehicle.current_driver = now_driver
        vehicle.save()
        VehicleDriverAssignment.objects.create(
            vehicle=vehicle, driver=then_driver, start_date=start.date(), end_date=start.date(),
            status=VehicleDriverAssignment.Status.ENDED,
        )
        report = trip_analytics.build_trip_report(
            user=UserFactory(role=_view_role()), vehicle_uuid=str(vehicle.uuid), start_at=start
        )
        assert report.record.driver_name == then_driver.get_full_name()


class TestPeriodReport:
    def test_rollups_match_the_trips(self):
        vehicle, _start = _seed_trip()
        report = trip_analytics.build_period_report(user=UserFactory(role=_view_role()), range_key="yesterday")
        assert report.summary["total_trips"] == 1
        assert report.summary["total_distance_km"] == 3.5
        assert report.summary["moving_seconds"] == 300
        assert report.summary["max_speed_kmh"] == 50.0
        assert [d["trips"] for d in report.day_rows] == [1]
        row = next(r for r in report.vehicle_rows if r["registration_number"] == vehicle.registration_number)
        assert row["trips"] == 1 and row["stops"] == 1  # pre-departure idling is outside the trip
        assert report.gps_total == 9  # a count for the summary card only
        assert not hasattr(report, "gps_rows")  # the fleet report never carries per-reading rows
        labels = [label for label, _v, _d in report.insights]
        assert "Maximum recorded speed" in labels and "Longest trip" in labels

    @pytest.mark.parametrize(
        "params, message",
        [
            ({"range": "custom", "from": "2026-01-10", "to": "2026-01-01"}, "on or before"),
            ({"range": "custom", "from": "2025-01-01", "to": "2025-06-01"}, "at most"),
            ({"range": "custom", "from": "", "to": ""}, "Select both"),
            ({"range": "custom", "from": "not-a-date", "to": "2026-01-01"}, "not valid"),
            ({"range": "decade"}, "Unknown date range"),
        ],
    )
    def test_invalid_ranges_are_rejected_with_a_message(self, params, message):
        with pytest.raises(trip_analytics.ReportError) as excinfo:
            trip_analytics.build_period_report(user=UserFactory(role=_view_role()), range_key=params.pop("range"),
                                               from_str=params.get("from", ""), to_str=params.get("to", ""))
        assert message in excinfo.value.message


class TestExportEndpoint:
    def test_requires_permission(self, client):
        client.force_login(UserFactory(role=None))
        assert client.get(reverse("trip-report-export"), {"type": "pdf"}).status_code == 403

    def test_pdf_download(self, client):
        _seed_trip()
        client.force_login(UserFactory(role=_view_role()))
        response = client.get(reverse("trip-report-export"), {"range": "yesterday", "type": "pdf"})
        assert response.status_code == 200
        assert response["Content-Type"] == "application/pdf"
        assert "attachment;" in response["Content-Disposition"]
        assert response.content.startswith(b"%PDF")
        assert AuditLog.objects.filter(action=AuditLog.Action.EXPORT, entity="TripReport").exists()

    def test_excel_download_has_the_analysis_sheets(self, client):
        from openpyxl import load_workbook

        _seed_trip()
        client.force_login(UserFactory(role=_view_role()))
        response = client.get(reverse("trip-report-export"), {"range": "last3", "type": "xlsx"})
        assert response.status_code == 200
        workbook = load_workbook(io.BytesIO(response.content))
        assert workbook.sheetnames == [
            "Summary", "Vehicle Analysis", "Daily Analysis", "Trip Details", "Stops", "Analysis",
        ]
        assert workbook["Trip Details"].max_row == 3  # header + 1 trip + totals row
        assert workbook["Daily Analysis"].max_row == 5  # header + 3 days + totals

    def test_fleet_report_has_no_gps_history(self, client, monkeypatch):
        """Fleet Trip Report = summary + analysis only — no per-reading GPS rows
        in either format (those live in the individual Trip Analysis Report)."""
        from openpyxl import load_workbook

        _seed_trip()
        client.force_login(UserFactory(role=_view_role()))
        workbook = load_workbook(io.BytesIO(
            client.get(reverse("trip-report-export"), {"range": "yesterday", "type": "xlsx"}).content
        ))
        assert "GPS History" not in workbook.sheetnames
        for sheet in workbook.worksheets:
            headers = {str(c.value) for c in sheet[1] if c.value}
            assert not headers & {"GPS timestamp", "Satellites", "Heading (°)", "GPS fix"}, sheet.title

        titles = _pdf_section_titles(monkeypatch)
        assert client.get(reverse("trip-report-export"), {"range": "yesterday", "type": "pdf"}).status_code == 200
        assert "4. Vehicle-wise Analysis" in titles and "6. Trip Details" in titles
        assert not [t for t in titles if "GPS History" in t]

    def test_no_data_is_a_readable_404(self, client):
        vehicle = VehicleFactory()
        TrackingDeviceFactory(vehicle=vehicle)
        client.force_login(UserFactory(role=_view_role()))
        response = client.get(reverse("trip-report-export"), {"range": "today", "type": "pdf"})
        assert response.status_code == 404
        assert "No trips" in response.json()["detail"]

    def test_bad_custom_range_is_a_400(self, client):
        client.force_login(UserFactory(role=_view_role()))
        response = client.get(
            reverse("trip-report-export"), {"range": "custom", "from": "2026-02-10", "to": "2026-02-01", "type": "pdf"}
        )
        assert response.status_code == 400
        assert "start date" in response.json()["detail"]

    def test_unknown_type_is_a_400(self, client):
        client.force_login(UserFactory(role=_view_role()))
        assert client.get(reverse("trip-report-export"), {"type": "docx"}).status_code == 400

    def test_client_user_cannot_export_another_clients_vehicle(self, client):
        mine, theirs = ClientFactory(), ClientFactory()
        other_vehicle, _ = _seed_trip(client_obj=theirs)
        client.force_login(UserFactory(role=_view_role(), client=mine))
        response = client.get(
            reverse("trip-report-export"), {"range": "yesterday", "vehicle": str(other_vehicle.uuid), "type": "xlsx"}
        )
        assert response.status_code == 404


class TestSingleTripEndpoints:
    def test_analysis_json_for_the_popup(self, client):
        vehicle, start = _seed_trip()
        client.force_login(UserFactory(role=_view_role()))
        response = client.get(reverse("trip-report-trip"), {"vehicle": str(vehicle.uuid), "start": start.isoformat()})
        assert response.status_code == 200
        body = response.json()
        assert body["trip_id"].startswith(vehicle.registration_number)
        assert body["distance_km"] == 3.5
        assert body["moving_seconds"] == 300 and body["idle_seconds"] == 120 and body["engine_off_seconds"] == 120
        assert body["max_speed_kmh"] == 50.0
        assert body["stop_count"] == 1 and body["longest_stop_seconds"] == 240
        assert body["gps_points"] == 8 and body["gps_fix_points"] == 8
        assert "points" not in body  # summary only — readings never go to the browser here

    @pytest.mark.parametrize("file_type, magic", [("pdf", b"%PDF"), ("xlsx", b"PK")])
    def test_trip_download(self, client, file_type, magic):
        vehicle, start = _seed_trip()
        client.force_login(UserFactory(role=_view_role()))
        response = client.get(
            reverse("trip-report-trip-export"), {"vehicle": str(vehicle.uuid), "start": start.isoformat(), "type": file_type}
        )
        assert response.status_code == 200
        assert response.content.startswith(magic)
        assert f"trip_{vehicle.registration_number}" in response["Content-Disposition"]

    def test_trip_excel_contains_every_reading_of_that_trip_only(self, client):
        from openpyxl import load_workbook

        vehicle, start = _seed_trip()
        # A second, later trip of the same vehicle must not leak into the first trip's file.
        for minute, ignition in [(0, True), (5, False)]:
            TelemetryEventFactory(device=vehicle.tracking_device, vehicle=vehicle, ignition=ignition, speed=20,
                                  timestamp=_yesterday_at(15) + datetime.timedelta(minutes=minute))
        client.force_login(UserFactory(role=_view_role()))
        response = client.get(
            reverse("trip-report-trip-export"), {"vehicle": str(vehicle.uuid), "start": start.isoformat(), "type": "xlsx"}
        )
        workbook = load_workbook(io.BytesIO(response.content))
        assert workbook.sheetnames == ["Trip Summary", "Trip Analysis", "GPS History"]
        gps = workbook["GPS History"]
        assert gps.max_row == 9  # header + all 8 readings of THIS trip (not the 15:00 trip, not pre-departure)
        headers = [c.value for c in gps[1]]
        for column in ("GPS timestamp", "Date", "Time", "Latitude", "Longitude", "Speed (km/h)", "Vehicle", "Driver",
                       "Location / address"):
            assert column in headers
        locations = [row[headers.index("Location / address")] for row in gps.iter_rows(min_row=2, values_only=True)]
        assert locations == [f"Point {m}" for m, *_ in READINGS[1:]]
        assert zipfile.ZipFile(io.BytesIO(response.content)).namelist().count("xl/media/image1.png") == 1  # route map

    def test_trip_pdf_includes_the_complete_gps_history(self, client, monkeypatch):
        vehicle, start = _seed_trip()
        titles = _pdf_section_titles(monkeypatch)
        client.force_login(UserFactory(role=_view_role()))
        response = client.get(
            reverse("trip-report-trip-export"), {"vehicle": str(vehicle.uuid), "start": start.isoformat(), "type": "pdf"}
        )
        assert response.status_code == 200
        assert "Complete GPS History" in titles
        assert {"Trip Summary", "Trip Route Map", "Trip Graphs", "Trip Analysis"} <= set(titles)

    def test_unknown_trip_start_is_a_404(self, client):
        vehicle, start = _seed_trip()
        client.force_login(UserFactory(role=_view_role()))
        response = client.get(
            reverse("trip-report-trip-export"),
            {"vehicle": str(vehicle.uuid), "start": (start + datetime.timedelta(seconds=30)).isoformat(), "type": "pdf"},
        )
        assert response.status_code == 404
        assert "could not be found" in response.json()["detail"]

    def test_missing_params_are_a_400(self, client):
        client.force_login(UserFactory(role=_view_role()))
        assert client.get(reverse("trip-report-trip"), {"vehicle": "x"}).status_code == 400

    def test_other_clients_trip_is_a_404(self, client):
        mine, theirs = ClientFactory(), ClientFactory()
        vehicle, start = _seed_trip(client_obj=theirs)
        client.force_login(UserFactory(role=_view_role(), client=mine))
        response = client.get(reverse("trip-report-trip"), {"vehicle": str(vehicle.uuid), "start": start.isoformat()})
        assert response.status_code == 404


def test_map_popup_has_downloads_but_no_open_vehicle_button(client):
    client.force_login(UserFactory(role=_view_role()))
    html = client.get(reverse("tracking:trip_report")).content.decode()
    assert "Open Vehicle" not in html and "tripRouteModalVehicleLink" not in html
    assert 'id="tripRouteModalMap"' in html and 'id="tripModalAnalysis"' in html
    assert 'data-trip-download="pdf"' in html and 'data-trip-download="xlsx"' in html


def test_list_api_exposes_the_same_trip_id(client):
    vehicle, start = _seed_trip()
    client.force_login(UserFactory(role=_view_role()))
    body = client.get(reverse("trip-report-data"), {"range": "yesterday"}).json()
    assert body["results"][0]["trip_id"] == f"{vehicle.registration_number}-{start:%Y%m%d-%H%M%S}"
    assert Decimal(body["results"][0]["distance_km"]) == Decimal("3.5")
